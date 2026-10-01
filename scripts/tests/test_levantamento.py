"""Teste de ponta a ponta do levantamento com uma pasta de cliente sintética.
Rodar: python scripts/tests/test_levantamento.py"""
import csv
import json
import os
import shutil
import sys
import tempfile

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), '..'))
import compras as C  # noqa: E402
import levantamento as LV  # noqa: E402


def _compra(chave, dia, ean, desc, qtd, vl, **k):
    r = {c: '' for c in C.CAMPOS_COMPRAS}
    r.update(chave_nfe=chave, numero='1', serie='1', data_emissao=dia, n_item='1', cnpj_emitente='01206820003031', nome_emitente='DIST',
             uf_emitente='SP', crt_emitente='3', cnpj_destinatario='33333333000133', uf_destinatario='SP', cod_produto='c' + chave[-2:],
             ean=ean, ean_trib=ean, descricao=desc, ncm='30049099', cest='1300100', cfop='5403', unid_comercial='UN', qtd=qtd, unid_trib='UN',
             qtd_trib=qtd, vl_merc=vl, cst='10', vbc_st=0, vbc_st_ret=0, p_icms_st=0, p_fcp_st=0, p_fcp_st_ret=0, p_st=0, tp_nf='1',
             fin_nfe='1', c_stat='100', situacao='Normal', arquivo_xml=chave + '.xml')
    r.update(k)
    return r


def _estoque(pos, ean, desc, ncm, qtd, valor, aliq, cest='1300100'):
    # CEST resolvido pelo EAN nas notas (como o intake faz): confirma o item da CAT 68 na triagem
    return dict(data_posicao=pos, arquivo=f'estoque {pos}.xls', ean=ean, cod_produto='1', descricao=desc, ncm=ncm, cest=cest,
                cest_origem='nfe_ean' if cest else 'nao_encontrado',
                qtd=qtd, unidade='UND', valor=valor, valor_total=round(qtd * valor, 2), grupo_pai='ETICO', cst_icms='00', totalizador=f'TC ({aliq})',
                aliq_totalizador=aliq, ean_dv_ok=True)


def _csv(caminho, campos, linhas):
    with open(caminho, 'w', newline='', encoding='utf-8-sig') as f:
        w = csv.DictWriter(f, fieldnames=campos, delimiter=';')
        w.writeheader()
        w.writerows(linhas)


def test_levantamento_sintetico():
    tmp = tempfile.mkdtemp()
    try:
        pasta = os.path.join(tmp, 'CLIENTE')
        os.makedirs(os.path.join(pasta, 'relatorio_parser'))
        compras = [
            _compra('A1', '2025-06-10', '7891000000001', 'PRODUTO X 10MG C/30 COMP', 4, 40.0, vbc_st=60.0, p_icms_st=18.0),
            _compra('A2', '2025-11-20', '7891000000001', 'PRODUTO X 10MG C/30 COMP', 3, 36.0, cst='60', vbc_st_ret=54.0, p_st=12.0),
            _compra('A3', '2026-02-01', '7891000000001', 'PRODUTO X 10MG C/30 COMP', 50, 500.0, vbc_st=900.0),   # após a vigência: ignorada
            _compra('B1', '2025-09-01', '7892000000002', 'PRODUTO Y 20MG C/30 COMP', 2, 30.0, ncm='33049990', cest='2000100'),
        ]
        _csv(os.path.join(pasta, 'relatorio_parser', 'compras.csv'), C.CAMPOS_COMPRAS, compras)
        est = [
            _estoque('2025-12-31', '7891000000001', 'PRODUTO X 10MG C/30 COMP', '30049099', 6.0, 11.0, 18.0),       # IX -> 01/01/2026
            _estoque('2025-12-31', '7893000000003', 'PRODUTO NUNCA COMPRADO 5MG', '30049099', 2.0, 9.0, 18.0),      # sem nota
            _estoque('2025-12-31', '7892000000002', 'PRODUTO Y 20MG C/30 COMP', '33049990', 2.0, 15.0, 18.0),       # XI -> 01/04/2026: falta posição
            _estoque('2025-12-31', '7894000000004', 'BRINQUEDO', '95030000', 1.0, 5.0, 18.0),                        # fora da lista
        ]
        campos = ['data_posicao', 'arquivo', 'ean', 'cod_produto', 'descricao', 'ncm', 'cest', 'cest_origem', 'qtd', 'unidade', 'valor',
                  'valor_total', 'grupo_pai', 'cst_icms', 'totalizador', 'aliq_totalizador', 'ean_dv_ok']
        _csv(os.path.join(pasta, 'relatorio_parser', 'estoque.csv'), campos, est)
        with open(os.path.join(pasta, 'cliente.json'), 'w', encoding='utf-8') as f:
            json.dump(dict(cliente='CLIENTE TESTE', cnpj='33333333000133', regime='SN', origem_regime='informado pelo usuário'), f)
        with open(os.path.join(pasta, 'relatorio_parser', 'intake_resumo.json'), 'w', encoding='utf-8') as f:
            json.dump(dict(estoque_posicoes_identicas=[[['estoque 31.12.2025.xls', '2025-12-31'], ['estoque 30.06.2026.xls', '2026-06-30']]]), f)

        r = LV.levantar(pasta)
        assert r['triados'] == 2                                                   # X e "nunca comprado"; Y não tem posição de 31/03
        from openpyxl import load_workbook
        wb = load_workbook(r['arquivo'])
        assert wb.sheetnames == ['Resumo', 'Anexo II (notas selecionadas)', 'Resumo por produto', 'Resumo por competência', 'Pendências']
        tr = load_workbook(os.path.join(pasta, 'relatorio_parser', 'trilha_levantamento.xlsx'))
        assert tr.sheetnames == ['Triagem estoque', 'Localização', 'Alocação de notas', 'Fora do crédito', 'Uso de notas', 'Premissas']
        tri = list(tr['Triagem estoque'].iter_rows(min_row=2, values_only=True))
        assert {x[1] for x in tri} == {'7891000000001', '7893000000003'}
        assert all(x[7] == 'IX' and x[10] == 'NCM' and x[12] == '01/01/2026' for x in tri)
        aloc = list(tr['Alocação de notas'].iter_rows(min_row=2, values_only=True))
        cab = [c.value for c in tr['Alocação de notas'][1]]
        ix = {n: i for i, n in enumerate(cab)}
        assert [a[ix['Chave NF-e']] for a in aloc] == ['A2', 'A1']                # mais recente primeiro; A3 ignorada (regra 4)
        a2, a1 = aloc
        assert a2[ix['Qtd usada do estoque']] == 3 and a1[ix['Qtd usada do estoque']] == 3        # 6 un.: 3 da A2 e 3 da A1
        assert abs(a2[ix['Base ST retido anteriormente unit.']] - 18.0) < 1e-6 and abs(a2[ix['Base ST x qtd usada']] - 54.0) < 1e-6
        assert abs(a1[ix['Base ST unit.']] - 15.0) < 1e-6 and abs(a1[ix['Base ST x qtd usada']] - 45.0) < 1e-6   # 60/4 x 3
        assert a2[ix['Alíquota ST']] == 12.0 and a1[ix['Alíquota ST']] == 18.0
        loc = list(tr['Localização'].iter_rows(min_row=2, values_only=True))
        sit = {x[1]: x[16] for x in loc}
        assert sit['7891000000001'] == 'Confirmado' and sit['7893000000003'] == 'Não localizado'
        pend = list(wb['Pendências'].iter_rows(min_row=2, values_only=True))
        tipos = [p[0] for p in pend]
        assert 'Posições de estoque idênticas' not in tipos and 'Falta posição de estoque' in tipos and 'Não localizado nas notas' in tipos
        assert not any(p[1] == 'BLOQUEANTE' for p in pend)                         # posição enviada vale como está
        faltam = [p for p in pend if p[0] == 'Falta posição de estoque']
        assert len(faltam) == 1 and '01/04/2026' in faltam[0][6]                   # só a data SEM arquivo (a de 01/01 tem posição)
        # crédito (Anexo V): A2 = (54,00 - 36,00) x 12% = 2,16; A1 = (45,00 - 30,00) x 18% = 2,70; mercadoria = 4,86
        a2s = list(wb['Anexo II (notas selecionadas)'].iter_rows(min_row=2, values_only=True))
        cab2 = [c.value for c in wb['Anexo II (notas selecionadas)'][1]]
        assert len(a2s) == 2 and [round(x[cab2.index('crédito')], 2) for x in a2s] == [2.16, 2.70]
        assert a2s[0][cab2.index('vl Mercd')] == 36.0 and a2s[0][cab2.index('BC ST')] == 54.0
        rp = {x[1]: x for x in wb['Resumo por produto'].iter_rows(min_row=2, values_only=True)}
        x = rp['7891000000001']
        assert x[7] == 4.86 and x[5] == 66.0 and x[6] == 99.0 and round(x[10], 4) == round(4.86 / 6, 4)        # itens 14, 15, 16 e 19
        assert rp['7893000000003'][12].startswith('Sem crédito')
        comp = list(wb['Resumo por competência'].iter_rows(min_row=2, values_only=True))
        assert [c[0] for c in comp] == ['02/2026'] and comp[0][2] == 4.86 and comp[0][3] == 'Simples Nacional'
        assert r['credito_total'] == 4.86 and r['anexo'] == 'V'
    finally:
        shutil.rmtree(tmp)



def _pasta(tmp, compras, est, regime='SN'):
    pasta = os.path.join(tmp, 'CLIENTE')
    os.makedirs(os.path.join(pasta, 'relatorio_parser'))
    _csv(os.path.join(pasta, 'relatorio_parser', 'compras.csv'), C.CAMPOS_COMPRAS, compras)
    campos = ['data_posicao', 'arquivo', 'ean', 'cod_produto', 'descricao', 'ncm', 'cest', 'cest_origem', 'qtd', 'unidade', 'valor',
              'valor_total', 'grupo_pai', 'cst_icms', 'totalizador', 'aliq_totalizador', 'ean_dv_ok']
    _csv(os.path.join(pasta, 'relatorio_parser', 'estoque.csv'), campos, est)
    with open(os.path.join(pasta, 'cliente.json'), 'w', encoding='utf-8') as f:
        json.dump(dict(cliente='CLIENTE TESTE', cnpj='33333333000133', regime=regime, origem_regime='informado'), f)
    with open(os.path.join(pasta, 'relatorio_parser', 'intake_resumo.json'), 'w', encoding='utf-8') as f:
        json.dump({}, f)
    return pasta


def test_item_de_nota_nunca_supre_dois_itens_do_estoque_e_ean_tem_prioridade():
    """O item só por descrição vem ANTES no arquivo, mas quem tem EAN escolhe a nota primeiro; a nota (4 un.) não é usada duas vezes."""
    tmp = tempfile.mkdtemp()
    try:
        compras = [_compra('N1', '2025-06-10', '7891000000001', 'PRODUTO X 10MG C/30 COMP', 4, 40.0, vbc_st=60.0, p_icms_st=18.0)]
        est = [_estoque('2025-12-31', '7899999999999', 'PRODUTO X 10MG C/30 COMP', '30049099', 4.0, 11.0, 18.0),   # sem EAN na nota
               _estoque('2025-12-31', '7891000000001', 'PRODUTO X 10MG C/30 COMP', '30049099', 4.0, 11.0, 18.0)]   # EAN da nota
        pasta = _pasta(tmp, compras, est)
        r = LV.levantar(pasta)
        from openpyxl import load_workbook
        tr = load_workbook(os.path.join(pasta, 'relatorio_parser', 'trilha_levantamento.xlsx'))
        cab = [c.value for c in tr['Alocação de notas'][1]]
        aloc = [dict(zip(cab, x)) for x in tr['Alocação de notas'].iter_rows(min_row=2, values_only=True)]
        assert [(a['EAN (estoque)'], a['Qtd usada do estoque']) for a in aloc] == [('7891000000001', 4)]
        sit = {x[1]: x[16] for x in tr['Localização'].iter_rows(min_row=2, values_only=True)}
        assert sit['7891000000001'] == 'Confirmado' and sit['7899999999999'] != 'Confirmado'
        assert r['credito_total'] == round((60.0 - 40.0) * 0.18, 2)
    finally:
        shutil.rmtree(tmp)


def test_localizacao_sem_prova_vai_para_fora_do_credito():
    tmp = tempfile.mkdtemp()
    try:
        compras = [_compra('N1', '2025-06-10', '7891000000001', 'LOSARTANA POTASSICA 50MG C/30 COMP EMS', 4, 40.0, vbc_st=60.0, p_icms_st=18.0),
                   _compra('N2', '2025-06-11', '7891000000002', 'LOSARTANA POTASSICA 50MG C/30 COMP MEDLEY', 4, 40.0, vbc_st=60.0, p_icms_st=18.0)]
        est = [_estoque('2025-12-31', '7899999999999', 'LOSARTANA POTASSICA 50MG C/30 COMP', '30049099', 2.0, 11.0, 18.0)]   # 2 laboratórios: ambíguo
        pasta = _pasta(tmp, compras, est)
        r = LV.levantar(pasta)
        assert r['credito_total'] == 0 and r['itens_fora_do_credito'] == 1 and r['fora_do_credito'] > 0
        from openpyxl import load_workbook
        wb = load_workbook(r['arquivo'])
        tipos = [p[0] for p in wb['Pendências'].iter_rows(min_row=2, values_only=True)]
        assert 'Fora do crédito por falta de prova' in tipos and 'Localização sujeita a validação' not in tipos
    finally:
        shutil.rmtree(tmp)


def test_rpa_anexo_iv_parcelas_e_relatorio():
    """Cliente RPA, fornecedor RPA com ST na nota (CST 10): BC ST x alíquota, definitivo; 12 parcelas somam o total;
    planilha e PDF falam do Anexo IV, não do Simples."""
    tmp = tempfile.mkdtemp()
    try:
        compras = [_compra('R1', '2025-06-10', '7891000000001', 'PRODUTO X 10MG C/30 COMP', 3, 30.0, cst='10', vbc_st=50.0, p_icms_st=18.0)]
        est = [_estoque('2025-12-31', '7891000000001', 'PRODUTO X 10MG C/30 COMP', '30049099', 3.0, 11.0, 18.0)]
        pasta = _pasta(tmp, compras, est, regime='RPA')
        r = LV.levantar(pasta)
        assert r['anexo'] == 'IV' and r['credito_total'] == 9.0                        # 50,00 x 18%
        from openpyxl import load_workbook
        wb = load_workbook(r['arquivo'])
        comp = [c for c in wb['Resumo por competência'].iter_rows(min_row=2, values_only=True)]
        assert len(comp) == 12 and comp[0][0] == '01/2026' and round(sum(c[2] for c in comp), 2) == 9.0
        resumo = ' '.join(str(c.value) for c in wb['Resumo']['A'] if c.value)
        assert 'RPA' in resumo and 'Simples Nacional, sem mudança' not in resumo and 'analogia' in resumo
        import gerar_relatorio as GR
        import fitz
        pdf = GR.gerar_pdf(pasta)
        texto = ' '.join(p.get_text() for p in fitz.open(pdf))
        assert 'Anexo IV' in texto and 'PGDAS' not in texto and 'Anexo V da Portaria' not in texto and 'analogia' in texto
    finally:
        shutil.rmtree(tmp)


def test_enquadramento_da_reducao_pelo_arquivo_e_leiame():
    """1ª rodada: redução em pRedBC sem enquadramento fica pendente e gera enquadramento_reducao.csv. 2ª rodada, com a coluna
    reducao preenchida: a linha é calculada. O LEIAME deixa de dizer "MODELO VAZIO"."""
    tmp = tempfile.mkdtemp()
    try:
        compras = [_compra('P1', '2025-06-10', '7891000000001', 'PRODUTO X 10MG C/30 COMP', 1, 34.52, cst='70', vbc_st=17.28, p_icms_st=18.0,
                           p_icms=18.0, p_red_bc=50.0)]                       # carga 9%: nenhuma regra automática decide
        est = [_estoque('2025-12-31', '7891000000001', 'PRODUTO X 10MG C/30 COMP', '30049099', 1.0, 11.0, 18.0)]
        pasta = _pasta(tmp, compras, est)
        with open(os.path.join(pasta, 'LEIAME.txt'), 'w', encoding='utf-8') as f:
            f.write('CLIENTE\n\nresultado_levantamento.xlsx    MODELO VAZIO\ncliente.json   dados\n')
        r = LV.levantar(pasta)
        assert r['credito_total'] == 0
        arq = os.path.join(pasta, LV.ENQUADRAMENTO)
        linhas = LV._ler_csv(arq)
        assert len(linhas) == 1 and linhas[0]['tipo'] == 'EAN' and linhas[0]['p_red_bc'] == '50' and linhas[0]['reducao'] == ''
        leiame = open(os.path.join(pasta, 'LEIAME.txt'), encoding='utf-8').read()
        assert 'MODELO VAZIO' not in leiame and 'levantamento concluído' in leiame and LV.ENQUADRAMENTO in leiame
        linhas[0]['reducao'], linhas[0]['justificativa'] = 'nao_aplicavel', 'RICMS/SP, teste'
        _csv(arq, LV.CAMPOS_ENQ, linhas)
        r2 = LV.levantar(pasta)
        assert r2['credito_total'] == round((17.28 - 34.52 * 0.5) * 0.18, 2)     # Anexo V, rótulo literal "não aplicável"
        assert LV._ler_csv(arq)[0]['justificativa'] == 'RICMS/SP, teste'                    # o arquivo preserva o que o usuário preencheu
        assert 'aguardando' not in ' '.join(str(c.value) for c in __import__('openpyxl').load_workbook(r2['arquivo'])['Resumo']['A'] if c.value)
    finally:
        shutil.rmtree(tmp)



def test_enquadramento_automatico_exige_identificacao_positiva():
    """Só o art. 3º, XXIV com o princípio ativo escrito na descrição é automático; marca, carga coincidente, NCM, art. 39 e
    redução de outra UF ficam para análise (casos do revisor fiscal)."""
    import enquadramento_auto as EA
    base = _compra('M1', '2025-06-10', '7891000000001', 'PARACETAMOL 750MG C 20 COMP', 1, 34.52, cst='70', vbc_st=17.28, p_icms_st=18.0,
                   p_icms=18.0, p_red_bc=61.11)
    ok = EA.classificar(base)
    assert ok['reducao'] == EA.APLICAVEL and ok['dispositivo'] == EA.DISPOSITIVO_XXIV and not ok['divergente'] and 'paracetamol' in ok['justificativa']
    assert EA.classificar(dict(base, descricao='TYLENOL 750MG C 10 COMP')) is None                 # marca não é identificação
    assert 'paracetamol' in EA.sugestao_marca('TYLENOL 750MG C 10 COMP')                            # ... vira sugestão para análise
    assert EA.classificar(dict(base, descricao='MEDICAMENTO DESCONHECIDO')) is None                 # carga 7% + NCM não bastam
    assert EA.classificar(dict(base, descricao='AMOXICILINA 500MG 21 CAPS')) is None                # amoxicilina sem clavulanato
    assert EA.classificar(dict(base, descricao='AMOXICILINA+CLAVULANATO 875+125MG'))['reducao'] == EA.APLICAVEL
    assert EA.classificar(dict(base, descricao='PARACETAMOL+CAFEINA 500+65MG')) is None              # associação fora do inciso
    assert EA.classificar(dict(base, descricao='TRAMADOL+PARACETAMOL 37,5+325MG'))['reducao'] == EA.APLICAVEL
    assert EA.classificar(dict(base, uf_emitente='PR')) is None                                     # inciso vale para operação interna SP
    assert EA.classificar(dict(base, data_emissao='2027-01-10')) is None                            # fora da vigência (31/12/2026)
    assert EA.classificar(dict(base, descricao='POTY FRUIT CAJU 200ML', ncm='22029900', p_red_bc=33.33)) is None   # art. 39: análise
    assert EA.classificar(dict(base, p_icms=12.0))['divergente']                                    # carga 4,67%: indício contrário


if __name__ == '__main__':
    for n, fn in list(globals().items()):
        if n.startswith('test_'):
            fn()
            print('ok', n)
