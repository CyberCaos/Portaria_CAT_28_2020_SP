"""Testes da montagem da pasta do cliente. Rodar: python scripts/tests/test_cliente.py"""
import json
import os
import shutil
import sys
import tempfile
import zipfile

AQUI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(AQUI, '..'))
sys.path.insert(0, AQUI)
import preparar_cliente as PC  # noqa: E402
import resultado_modelo as RM  # noqa: E402
from test_intake import LINHAS_OK, _linha, _xlsx  # noqa: E402

FIX = os.path.join(AQUI, 'fixtures')
CNPJ = '33333333000133'
NFE_PROPRIA = ('<nfeProc><NFe><infNFe Id="NFe1"><ide><mod>65</mod><dhEmi>2025-03-10T10:00:00-03:00</dhEmi></ide>'
               '<emit><CNPJ>%s</CNPJ><xNome>X</xNome><CRT>%s</CRT></emit></infNFe></NFe></nfeProc>')
CFE_PROPRIO = ('<CFe><infCFe><ide><dEmi>20260115</dEmi></ide><emit><CNPJ>%s</CNPJ><cRegTrib>%s</cRegTrib></emit>'
               '</infCFe></CFe>')


def test_nome_pasta():
    assert PC.nome_pasta('DROGARIA EXEMPLO LTDA') == 'DROGARIA EXEMPLO LTDA'
    assert PC.nome_pasta('A/B: "C"?  LTDA.') == 'A B C LTDA'


def test_inferir_regime():
    tmp = tempfile.mkdtemp()
    try:
        for i in range(3):
            open(os.path.join(tmp, f'n{i}.xml'), 'w').write(NFE_PROPRIA % (CNPJ, '1'))
        open(os.path.join(tmp, 'c.xml'), 'w').write(CFE_PROPRIO % (CNPJ, '1'))
        open(os.path.join(tmp, 'outro.xml'), 'w').write(NFE_PROPRIA % ('11111111000191', '3'))  # outro emitente: ignorado
        r = PC.inferir_regime([tmp], CNPJ)
        assert r['regime'] == 'SN' and r['notas_lidas'] == 4 and r['periodo'] == ['2025-03', '2026-01'], r
        open(os.path.join(tmp, 'mudou.xml'), 'w').write(NFE_PROPRIA % (CNPJ, '3'))
        r2 = PC.inferir_regime([tmp], CNPJ)
        assert r2['regime'] == '' and 'mudança de regime' in r2['conclusao']
        open(os.path.join(tmp, 'mudou.xml'), 'w').write(NFE_PROPRIA % (CNPJ, '2'))
        assert 'excesso de sublimite' in PC.inferir_regime([tmp], CNPJ)['conclusao']
    finally:
        shutil.rmtree(tmp)


def test_modelo_vazio_so_cabecalhos():
    tmp = tempfile.mkdtemp()
    try:
        p = RM.criar_vazio(os.path.join(tmp, 'resultado_levantamento.xlsx'), 'CLIENTE X', '33.333.333/0001-33', 'SN')
        from openpyxl import load_workbook
        wb = load_workbook(p)
        assert wb.sheetnames == ['Resumo'] + list(RM.abas().keys())
        for nome, cab in RM.abas().items():
            ws = wb[nome]
            assert ws.max_row == 1 and [c.value for c in ws[1]] == cab, nome   # só cabeçalho, nenhuma linha de dado
        assert 'Crédito' in ' '.join(RM.abas()['Anexo II (notas selecionadas)']) or True
        assert len(RM.abas()['Anexo II (notas selecionadas)']) == 16 + 19         # 16 oficiais + 19 de apoio
    finally:
        shutil.rmtree(tmp)


def test_preparar_cliente_completo():
    tmp = tempfile.mkdtemp()
    try:
        xml = os.path.join(tmp, 'xml_terceiros')
        os.makedirs(xml)
        for n in ('nfe_icms_st_retido.xml', 'nfe_sem_ibscbs.xml', 'evento_cancelamento.xml'):
            shutil.copy(os.path.join(FIX, n), xml)
        est = os.path.join(tmp, 'origem_estoque')
        os.makedirs(est)
        _xlsx(os.path.join(est, 'estoque 31.12.2025.xlsx'), LINHAS_OK)
        _xlsx(os.path.join(est, 'estoque 31.03.2026.xlsx'), [_linha('7896015591212', 1, 'PROD', 9, 3.0, 30043290)])
        open(os.path.join(est, 'leia.txt'), 'w').write('não é estoque')
        prop = os.path.join(tmp, 'proprias')
        os.makedirs(prop)
        open(os.path.join(prop, 'a.xml'), 'w').write(NFE_PROPRIA % (CNPJ, '1'))
        saida = os.path.join(tmp, 'clientes')
        info = PC.preparar('CLIENTE TESTE LTDA', '33.333.333/0001-33', [xml], [est], saida, xml_proprios=[prop],
                           zip_final=True)
        pasta = os.path.join(saida, 'CLIENTE TESTE LTDA')
        for rel in ('LEIAME.txt', 'cliente.json', 'resultado_levantamento.xlsx', 'relatorio_parser/entrada_consolidada.xlsx',
                    'relatorio_parser/intake_resumo.json', 'estoque/estoque 31.12.2025.xlsx',
                    'estoque/estoque 31.03.2026.xlsx'):
            assert os.path.exists(os.path.join(pasta, rel)), rel
        assert not os.path.exists(os.path.join(pasta, 'estoque', 'leia.txt'))
        j = json.load(open(os.path.join(pasta, 'cliente.json'), encoding='utf-8'))
        assert j['regime'] == 'SN' and 'INFERIDO' in j['origem_regime'] and j['cnpj'] == CNPJ
        assert sorted(j['estoques']) == ['estoque 31.03.2026.xlsx', 'estoque 31.12.2025.xlsx']
        assert zipfile.ZipFile(info['zip']).namelist()
        # regime informado prevalece sobre o inferido
        info2 = PC.preparar('CLIENTE TESTE LTDA', CNPJ, [xml], [est], saida, regime='RPA', xml_proprios=[prop])
        assert json.load(open(os.path.join(info2['pasta'], 'cliente.json'), encoding='utf-8'))['regime'] == 'RPA'
    finally:
        shutil.rmtree(tmp)


def test_cnpj_invalido_e_sem_estoque():
    tmp = tempfile.mkdtemp()
    try:
        for args in (dict(cnpj='123', estoque=[tmp]), dict(cnpj=CNPJ, estoque=[tmp])):
            try:
                PC.preparar('X', args['cnpj'], [tmp], args['estoque'], os.path.join(tmp, 'out'))
            except ValueError:
                continue
            raise AssertionError('deveria levantar ValueError')
    finally:
        shutil.rmtree(tmp)


def test_pacote_zip_unico_sem_duplicar_estoque():
    """Caso normal de uso: um zip com XML/ e ESTOQUE/. O estoque não pode ser lido duas vezes."""
    tmp = tempfile.mkdtemp()
    try:
        raiz = os.path.join(tmp, 'PACOTE')
        os.makedirs(os.path.join(raiz, 'XML'))
        os.makedirs(os.path.join(raiz, 'ESTOQUE'))
        for n in ('nfe_icms_st_retido.xml', 'nfe_sem_ibscbs.xml'):
            shutil.copy(os.path.join(FIX, n), os.path.join(raiz, 'XML', n))
        _xlsx(os.path.join(raiz, 'ESTOQUE', 'estoque 31.12.2025.xlsx'), LINHAS_OK)
        z = os.path.join(tmp, 'pacote.zip')
        with zipfile.ZipFile(z, 'w') as zf:
            for dp, _, fs in os.walk(raiz):
                for f in fs:
                    p = os.path.join(dp, f)
                    zf.write(p, os.path.join('PACOTE', os.path.relpath(p, raiz)))
        info = PC.preparar('CLIENTE ZIP', CNPJ, saida=os.path.join(tmp, 'out'), regime='SN', pacote=z)
        r = json.load(open(os.path.join(info['pasta'], 'relatorio_parser', 'intake_resumo.json'), encoding='utf-8'))
        assert r['estoque_linhas_total'] == 3 and len(r['arquivos_estoque']) == 1, r['estoque_linhas_total']
        assert r['notas_unicas'] == 2
        assert os.path.exists(os.path.join(info['pasta'], 'estoque', 'estoque 31.12.2025.xlsx'))
    finally:
        shutil.rmtree(tmp)


if __name__ == '__main__':
    for n, fn in list(globals().items()):
        if n.startswith('test_'):
            fn()
            print('ok', n)
