"""Testes da etapa de entrada (pacote zip/pasta -> entrada_consolidada.xlsx) e do leitor de estoque.
Fixtures sintéticas que imitam o layout real do relatório "Posição de Estoque (Inventário)" (sem dados de cliente).
Rodar: python scripts/tests/test_intake.py"""
import os
import shutil
import sys
import tempfile
import zipfile
from datetime import date

AQUI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(AQUI, '..'))
import estoque as E  # noqa: E402
import intake  # noqa: E402
import triagem as PE  # noqa: E402

FIX = os.path.join(AQUI, 'fixtures')
CAB = ['Código de Barras', None, 'Produto ID', 'Descrição do Produto', None, None, 'Grupo pai', 'Grupo filho', 'Qtde.',
       'Preço Custo Médio', 'Total Preço Custo Médio', 'Preço Venda', 'Total Preço Venda', 'Unidade', 'CST', 'NCM',
       'CST PIS', 'CST COFINS', 'Totalizador']


def _linha(ean, pid, desc, qtd, custo, ncm, tot='TC (18.00)', cst='00'):
    total = None if qtd is None or custo is None else round(qtd * custo, 2)
    return [ean, None, pid, desc, None, None, 'ETICO', 'ETICO', qtd, custo, total,
            custo * 1.3 if custo else None, None, 'UND', cst, ncm, '04', '04', tot]


def _xlsx(caminho, linhas):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append([None] * 19)
    ws.append(['Relatório', None, None, None, 'Posição de Estoque (Inventário)'])
    ws.append([None] * 19)
    ws.append(CAB)
    for l in linhas:
        ws.append(l)
    total = sum(l[10] or 0 for l in linhas)
    ws.append([None] * 9 + [0, round(total, 2)])   # rodapé: sem produto, só totais
    wb.save(caminho)


LINHAS_OK = [
    _linha('7896015591212', 46300, 'PRODUTO A 30 COMP', 2, 10.5, 30043290),
    _linha('7891317038878', 46789, 'PRODUTO B 10MG', 3, 4.0, 30049049),
    _linha('7897074002015', 12886, 'ANTIGRIPAL', 1, 2.0, 9021000),   # NCM com 7 dígitos (Excel perdeu o zero)
]


def _pacote(tmp, linhas_estoque=LINHAS_OK, datas=('31.12.2025', '31.03.2026')):
    raiz = os.path.join(tmp, 'pacote')
    os.makedirs(os.path.join(raiz, 'compras'))
    os.makedirs(os.path.join(raiz, 'estoque'))
    for n in ('nfe_icms_st_retido.xml', 'nfe_para_cancelar.xml', 'evento_cancelamento.xml', 'nfe_sem_ibscbs.xml'):
        shutil.copy(os.path.join(FIX, n), os.path.join(raiz, 'compras', n))
    open(os.path.join(raiz, 'compras', 'quebrado.xml'), 'w').write('<nfeProc>')
    for d in datas:
        _xlsx(os.path.join(raiz, 'estoque', f'estoque {d}.xlsx'), linhas_estoque)
    return raiz


def test_data_do_nome_num_br_e_helpers():
    assert E.data_do_nome('estoque 31.12.2025.xlsx') == date(2025, 12, 31)
    assert E.data_do_nome('ESTOQUE_01-04-2026.csv') == date(2026, 4, 1)
    assert E.data_do_nome('estoque.xlsx') is None
    assert E.num_br('1.234,56') == 1234.56 and E.num_br('1,5') == 1.5 and E.num_br('1,234.5') == 1234.5
    assert E.num_br('') is None and E.num_br('abc') is None and E.num_br(float('nan')) is None
    assert E.digitos(9021000.0) == '9021000' and E.digitos(7896015591212) == '7896015591212' and E.digitos(None) == ''
    assert E.normalizar_ncm(9021000) == ('09021000', 'NCM com 7 dígitos: zero à esquerda restaurado')
    assert E.normalizar_ncm('3004.90.99') == ('30049099', None)
    assert E.normalizar_ncm('123')[0] == '123' and E.normalizar_ncm('123')[1].startswith('NCM inválido')
    assert E.gtin_valido('7896015591212') and not E.gtin_valido('7896015591213') and not E.gtin_valido('ABC')
    assert E.aliquota_totalizador('TC (18.00)') == 18.0 and E.aliquota_totalizador('F (0.00)') == 0.0


def test_layout_real_obrigatorios_e_rodape():
    tmp = tempfile.mkdtemp()
    try:
        p = os.path.join(tmp, 'estoque 31.12.2025.xlsx')
        _xlsx(p, LINHAS_OK)
        r = E.ler(p, date(2025, 12, 31))
        assert r['bloqueantes'] == [] and len(r['linhas']) == 3 and r['rejeitadas'] == []
        assert r['colunas']['ean'] == 'Código de Barras' and r['colunas']['valor'] == 'Preço Custo Médio'
        assert r['colunas']['valor_total'] == 'Total Preço Custo Médio' and r['colunas']['cst_icms'] == 'CST'
        l = r['linhas'][2]
        assert l['ncm'] == '09021000' and l['ean'] == '7897074002015' and l['valor'] == 2.0
        assert r['linhas'][0]['aliq_totalizador'] == 18.0 and r['linhas'][0]['cest'] == ''
        assert r['total_rodape'] == round(21 + 12 + 2, 2)
        assert any('zero à esquerda' in a for a in r['avisos']) and any('sem CEST' in a for a in r['avisos'])
    finally:
        shutil.rmtree(tmp)


def test_linhas_rejeitadas_por_campo_obrigatorio():
    tmp = tempfile.mkdtemp()
    try:
        p = os.path.join(tmp, 'estoque 31.12.2025.xlsx')
        ruins = [
            _linha(None, 1, 'SEM EAN', 1, 5.0, 30049099),
            _linha('SEM GTIN', 2, 'EAN TEXTO', 1, 5.0, 30049099),
            _linha('12345', 8, 'EAN CURTO', 1, 5.0, 30049099),
            _linha('7896015591212', 3, '', 1, 5.0, 30049099),
            _linha('7891317038878', 4, 'NCM CURTO', 1, 5.0, 123),
            _linha('7897074002015', 5, 'SEM VALOR', 1, None, 30049099),
            _linha('7896094210967', 6, 'VALOR ZERO', 1, 0, 30049099),
            _linha('7898040321499', 7, 'OK', 1, 5.0, 30049099),
        ]
        _xlsx(p, ruins)
        r = E.ler(p, date(2025, 12, 31))
        assert [x['descricao'] for x in r['linhas']] == ['OK']
        motivos = ' | '.join(x['motivo'] for x in r['rejeitadas'])
        for esperado in ('EAN ausente', 'EAN sem dígitos', 'EAN com tamanho inválido', 'Descrição ausente', 'NCM inválido', 'Valor ausente',
                         'Valor zero/negativo'):
            assert esperado in motivos, (esperado, motivos)
        assert len(r['rejeitadas']) == 7 and any('rejeitada' in a for a in r['avisos'])
    finally:
        shutil.rmtree(tmp)


def test_coluna_obrigatoria_ausente_e_bloqueante():
    tmp = tempfile.mkdtemp()
    try:
        p = os.path.join(tmp, 'estoque 01.01.2026.csv')
        open(p, 'w', encoding='utf-8').write('cod;descricao;ncm;quantidade;valor\nA1;PRODUTO;30049099;1;5,0\n')
        r = E.ler(p, date(2026, 1, 1))
        assert r['linhas'] == [] and any('EAN' in b for b in r['bloqueantes']), r['bloqueantes']
        q = os.path.join(tmp, 'estoque 02.01.2026.csv')
        open(q, 'w').write('a;b;c\n1;2;3\n')
        assert 'Cabeçalho não reconhecido' in E.ler(q, date(2026, 1, 2))['bloqueantes'][0]
    finally:
        shutil.rmtree(tmp)


def test_valor_derivado_do_total_quando_so_ha_total():
    tmp = tempfile.mkdtemp()
    try:
        p = os.path.join(tmp, 'estoque 01.01.2026.csv')
        open(p, 'w', encoding='utf-8').write(
            'Código de Barras;Produto ID;Descrição do Produto;Qtde.;Total Preço Custo Médio;NCM\n'
            '7896015591212;1;PRODUTO;4;20,00;30049099\n')
        r = E.ler(p, date(2026, 1, 1))
        assert r['bloqueantes'] == [] and r['linhas'][0]['valor'] == 5.0
    finally:
        shutil.rmtree(tmp)


def test_resolver_cest_por_ean():
    compras = [dict(ean='7896015591212', cest='1300100'), dict(ean='7896015591212', cest='1300100'),
               dict(ean='7891317038878', cest='1300100'), dict(ean='7891317038878', cest='1300200'),
               dict(ean='SEM GTIN', cest='1300100'), dict(ean='7897074002015', cest='')]
    est = [dict(ean='7896015591212', cest=''), dict(ean='7891317038878', cest=''), dict(ean='7897074002015', cest=''),
           dict(ean='7898040321499', cest='2300100')]
    cob = PE.resolver_cest(est, PE.indice_cest_por_ean(compras))
    assert [e['cest_origem'] for e in est] == ['nfe_ean', 'nfe_ean_ambiguo', 'nao_encontrado', 'estoque']
    assert est[0]['cest'] == '1300100' and est[2]['cest'] == '' and cob['total'] == 4


def test_pacote_pasta_e_zip_com_posicoes_identicas():
    tmp = tempfile.mkdtemp()
    try:
        raiz = _pacote(tmp)
        z = os.path.join(tmp, 'pacote.zip')
        with zipfile.ZipFile(z, 'w') as zf:
            for dp, _, fs in os.walk(raiz):
                for f in fs:
                    p = os.path.join(dp, f)
                    zf.write(p, os.path.relpath(p, raiz))
        for origem in (raiz, z):
            saida = os.path.join(tmp, 'saida_' + os.path.basename(origem))
            r = intake.consolidar(origem, saida, cnpj='99.999.999/0001-91')
            assert r['xml_com_erro'] == 1 and r['eventos_cancelamento'] == 1
            assert r['notas_unicas'] == 3 and r['itens_de_compra'] == 3 and r['itens_com_bc_st_retida'] == 1
            assert sorted(a['data_posicao'] for a in r['arquivos_estoque']) == ['2025-12-31', '2026-03-31']
            assert r['estoque_linhas_total'] == 6 and r['estoque_linhas_rejeitadas'] == 0
            assert r['papel_empresa'].get('entrada') == 2 and r['papel_empresa'].get('terceiros') == 1
            assert len(r['estoque_posicoes_identicas']) == 1
            assert any(a.startswith('Info') and 'mesmo conteúdo' in a for a in r['alertas']), r['alertas']   # vale como enviada
            for n in ('entrada_consolidada.xlsx', 'compras.csv', 'estoque.csv', 'intake_resumo.json'):
                assert os.path.exists(os.path.join(saida, n)), n
        from openpyxl import load_workbook
        wb = load_workbook(os.path.join(tmp, 'saida_pacote', 'entrada_consolidada.xlsx'))
        assert wb.sheetnames[:4] == ['Resumo', 'Compras', 'Estoque', 'Estoque rejeitadas']
    finally:
        shutil.rmtree(tmp)


def test_posicoes_diferentes_nao_alertam_identicas():
    tmp = tempfile.mkdtemp()
    try:
        raiz = _pacote(tmp, datas=('31.12.2025',))
        _xlsx(os.path.join(raiz, 'estoque', 'estoque 31.03.2026.xlsx'),
              [_linha('7896015591212', 46300, 'PRODUTO A 30 COMP', 5, 10.5, 30043290)])
        r = intake.consolidar(raiz, os.path.join(tmp, 'saida'))
        assert r['estoque_posicoes_identicas'] == []
        assert not any('mesmo conteúdo' in a for a in r['alertas'])
    finally:
        shutil.rmtree(tmp)


if __name__ == '__main__':
    for n, fn in list(globals().items()):
        if n.startswith('test_'):
            fn()
            print('ok', n)
