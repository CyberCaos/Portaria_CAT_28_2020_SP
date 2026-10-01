"""Modelo da planilha `resultado_levantamento.xlsx` — uma por cliente.

`criar_vazio` gera só estrutura (abas, cabeçalhos, formatação), usada por `preparar_cliente.py`. `levantamento.py` grava um
arquivo novo com todas as abas do cliente (`abas`) preenchidas e, à parte, a trilha técnica (`abas_trilha`) em
relatorio_parser/trilha_levantamento.xlsx.
"""
import json
import os
from datetime import date

RAIZ = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))

TRIAGEM = ['Data da posição', 'EAN', 'Descrição', 'NCM', 'Qtd em estoque', 'Valor unitário (custo médio)', 'Valor total',
           'Anexo CAT 68', 'Segmento', 'Item CAT 68', 'Critério', 'CEST confirma', 'Data de revogação', 'CEST', 'Origem do CEST', 'Ato revogador',
           'Sinalizações']
LOCALIZACAO = ['Data da posição', 'EAN', 'Descrição (estoque)', 'Qtd em estoque', 'Método', 'Nível', 'Produto na nota (descrição)',
               'EAN na nota', 'Linhas de nota elegíveis', 'Qtd disponível nas notas (un. estoque)', 'Qtd coberta', 'Qtd descoberta',
               'Cobertura', 'Valor do estoque coberto', 'Fator (predominante)', 'Confiança do fator', 'Situação', 'Observações']
ALOCACAO = ['Data da posição', 'EAN (estoque)', 'Descrição (estoque)', 'Data de revogação', 'Origem da localização',
            'Chave NF-e', 'Item', 'Emissão', 'Descrição na nota', 'EAN na nota', 'Un. nota', 'Qtd nota', 'Fator', 'Origem do fator',
            'Confiança do fator', 'Obs. do fator', 'Qtd da linha em un. estoque', 'Qtd usada do estoque', 'Vl mercadoria unit.',
            'Base ST unit.', 'Base ST retido anteriormente unit.', 'Base ST total unit.', 'Base ST x qtd usada', 'Alíquota ST',
            'Origem da alíquota', 'CST/CSOSN', 'CRT emitente', 'UF emitente', 'UF destinatário', 'CFOP', 'pRedBCST', 'Sem base de ST']
FORA_DO_CREDITO = ['Data da posição', 'EAN (estoque)', 'Descrição (estoque)', 'Data de revogação', 'Motivo', 'Chave NF-e', 'Item', 'Emissão',
                   'Descrição na nota', 'Qtd do estoque (hipótese)', 'Fator', 'VlMerc líquido', 'BC ST', 'Alíquota', 'Crédito potencial da linha',
                   'Crédito potencial do item']
USO_NOTAS = ['Chave NF-e', 'Item', 'Descrição na nota', 'Qtd comprada', 'Qtd usada (todas as posições)', 'Saldo', 'Nº de alocações',
             'Itens do estoque atendidos']
PENDENCIAS = ['Tipo', 'Severidade', 'Data da posição', 'EAN', 'Descrição', 'Valor em estoque', 'Detalhe', 'Ação sugerida']
PREMISSAS = ['Premissa', 'Origem', 'Data', 'Responsável']


def _colunas_anexo_ii():
    with open(os.path.join(RAIZ, 'assets', 'anexo-ii-colunas.json'), encoding='utf-8') as f:
        j = json.load(f)
    return [c['rotulo'] for c in j['colunas_tabela_a']]


def abas(parser_cols=None):
    """Planilha do CLIENTE (resultado_levantamento.xlsx), depois da aba 'Resumo': só o que ele usa para lançar e conferir."""
    return {
        'Anexo II (notas selecionadas)': _colunas_anexo_ii() + ['Fórmula aplicada', 'Linha da tabela (Anexo IV/V)', 'Alertas', 'EAN',
                                                              'Descrição (estoque)', 'Data de revogação', 'Vl mercadoria bruto (vProd)', 'Desconto (vDesc)',
                                                              'Crédito com VlMerc = vProd (alternativa)',
                                                              'Observações do cálculo', 'Origem do pRedBc usado', 'pRedBCST da nota',
                                                              'pRedBC da nota', 'Enquadramento da redução', 'Valores alternativos', 'Origem da alíquota',
                                                              'Evidência de entrada', 'Revisar entrada (prioridade)', 'Evidências da identidade (CAT 68)'],
        'Resumo por produto': ['Cód. produto', 'EAN', 'Descrição', 'Data de revogação', 'Qtd estoque', 'Total vl Mercd (item 14)',
                               'Total BC ST (item 15)', 'Total crédito (item 16)', 'Unit. vl Mercd (item 17)', 'Unit. BC ST (item 18)',
                               'Unit. crédito (item 19)', 'Qtd sem nota (descoberto)', 'Situação do crédito',
                               'Observações do cálculo'],
        'Resumo por competência': ['Competência', 'Parcela', 'Crédito a lançar', 'Regime', 'Forma de lançamento', 'Observação'],
        'Pendências': PENDENCIAS,
    }


def abas_trilha():
    """Trilha técnica (relatorio_parser/trilha_levantamento.xlsx): passos intermediários, auditoria e insumo do PDF."""
    return {
        'Triagem estoque': TRIAGEM,
        'Localização': LOCALIZACAO,
        'Alocação de notas': ALOCACAO,
        'Fora do crédito': FORA_DO_CREDITO,
        'Uso de notas': USO_NOTAS,
        'Premissas': PREMISSAS,
    }


def estilizar(ws, cab, larguras=None, altura=32):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    fill = PatternFill('solid', fgColor='1F3864')
    for j, c in enumerate(ws[1], 1):
        c.fill, c.font = fill, Font(bold=True, color='FFFFFF')
        c.alignment = Alignment(wrap_text=True, vertical='center')
        ws.column_dimensions[get_column_letter(j)].width = (larguras or {}).get(cab[j - 1], max(12, min(len(str(cab[j - 1])) + 4, 38)))
    ws.row_dimensions[1].height = altura
    ws.freeze_panes = 'A2'


def leia_me(ws, cliente, cnpj, regime, status, linhas_extra=()):
    from openpyxl.styles import Font
    linhas = [
        ('RESULTADO DO LEVANTAMENTO', True),
        (f'Cliente: {cliente}', False), (f'CNPJ: {cnpj}', False), (f'Regime: {regime or "não informado"}', False),
        (f'Gerado em {date.today().strftime("%d/%m/%Y")} pela skill credito-icms-estoque-st.', False),
        (status, False),
        ('Base legal: Portaria CAT 28/2020 (Anexos I a V) e Portaria CAT 68/2019 (itens revogados).', False),
        *[(t, False) for t in linhas_extra],
    ]
    for t, negrito in linhas:
        ws.append([t])
        ws.cell(ws.max_row, 1).font = Font(bold=negrito, size=13 if negrito else 11)
    ws.column_dimensions['A'].width = 140


def criar_vazio(caminho: str, cliente: str, cnpj: str, regime: str = '') -> str:
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = 'Resumo'
    leia_me(ws, cliente, cnpj, regime, 'Status: modelo sem dados. As abas só têm cabeçalhos; o levantamento as preenche '
                                       '(python scripts/levantamento.py <pasta do cliente>).')
    for nome, cab in abas().items():
        a = wb.create_sheet(nome[:31])
        a.append(cab)
        estilizar(a, cab)
    os.makedirs(os.path.dirname(os.path.abspath(caminho)), exist_ok=True)
    wb.save(caminho)
    return caminho
