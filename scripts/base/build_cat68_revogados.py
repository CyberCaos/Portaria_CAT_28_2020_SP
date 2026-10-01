"""Gera a planilha de produtos revogados da Portaria CAT 68/2019 (base para levantamento de crédito - CAT 28/2020)."""
import re, sys, os, io, csv, contextlib
from collections import Counter
from datetime import date

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
with contextlib.redirect_stdout(io.StringIO()):
    from parse_cat68 import out as ROWS, parse_rev, BASE

REF = date(2026, 9, 30)  # data de referência da extração

# Atos revogadores: (data do ato, DOE, início da vigência, fonte da vigência) — conferidos nas páginas de cada ato
ACTS = {
    'CAT-57/21': (date(2021, 8, 6), date(2021, 8, 7), date(2021, 8, 7), 'Art. 4º do ato: vigor na data da publicação (DOE)'),
    'CAT-91/21': (date(2021, 12, 23), date(2021, 12, 24), date(2022, 1, 1), 'Art. 4º do ato'),
    'SRE-18/22': (date(2022, 3, 29), date(2022, 3, 30), date(2022, 4, 1), 'Art. 3º do ato'),
    'SRE-69/22': (date(2022, 9, 14), date(2022, 9, 15), date(2022, 10, 1), 'Art. 4º do ato'),
    'SRE-64/25': (date(2025, 10, 1), date(2025, 10, 2), date(2026, 1, 1), 'Art. 1º do ato: "a partir de 1º de janeiro de 2026"'),
    'SRE-94/25': (date(2025, 12, 22), date(2025, 12, 23), date(2026, 4, 1), 'Art. 4º do ato'),
    'SRE-09/26': (date(2026, 3, 17), date(2026, 3, 18), date(2026, 7, 1), 'Art. 3º do ato'),
    'SRE-19/26': (date(2026, 4, 29), date(2026, 4, 30), date(2026, 8, 1), 'Art. 3º do ato'),
    'SRE-20/26': (date(2026, 5, 4), date(2026, 5, 5), date(2026, 8, 1), 'Art. 3º do ato'),
    'SRE-34/26': (date(2026, 6, 29), date(2026, 6, 30), date(2026, 10, 1), 'Art. 3º do ato'),
}
ORDEM = {k: i for i, k in enumerate('I II III IV V VI VII VIII IX X XI XII XIII XIV XV XVI XVII XVIII XIX XX XXI XXII'.split())}


def sortnum(item):
    return [int(x) for x in item.split('.')]


def situacao(d):
    if d['vig'] <= REF:
        return 'Revogado (em vigor)'
    return 'Revogação futura (não vigente em %s)' % REF.strftime('%d/%m/%Y')


# A 1ª linha de cada (anexo,item) é a redação/estado corrente; as demais são redações anteriores
cur = {}
for r in ROWS:
    cur.setdefault((r['anexo'], r['item']), r)

itens, annexes = [], {}
for (ax, it), r in cur.items():
    own = 'revogad' in re.sub(r'\s+', '', r['first_cell']).lower()
    arev = r['anexo_rev']
    if own:
        x = parse_rev(r['first_cell'])
        tipo = 'Item revogado'
    elif arev:
        x = arev
        tipo = 'Anexo inteiro revogado'
    else:
        continue
    ato = x['ato'].replace('SRE-9/26', 'SRE-09/26')
    d_ato, d_doe, d_vig, _ = ACTS[ato]
    if x['vigencia']:  # confere data lida na página com a do ato
        dd, mm, yy = x['vigencia'].split('-')
        assert date(int(yy), int(mm), int(dd)) == d_vig, (ax, it, x, d_vig)
    itens.append(dict(anexo=ax, anexo_titulo=r['anexo_titulo'], item=it, cest=r['cest'], ncm=r['ncm'],
                      descricao=r['descricao'], tipo=tipo, ato=ato, data_ato=d_ato, doe=d_doe, vig=d_vig))
    if arev:
        annexes.setdefault(ax, (r['anexo_titulo'], arev['ato'].replace('SRE-9/26', 'SRE-09/26')))

itens.sort(key=lambda d: (ORDEM[d['anexo']], sortnum(d['item'])))

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

wb = Workbook()
hdrfill = PatternFill('solid', fgColor='1F3864')
hdrfont = Font(bold=True, color='FFFFFF')


def sheet(ws, headers, data, widths, datecols=()):
    ws.append(headers)
    for c in ws[1]:
        c.fill = hdrfill
        c.font = hdrfont
        c.alignment = Alignment(wrap_text=True, vertical='center')
    for row in data:
        ws.append(row)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for col in datecols:
        for c in ws[get_column_letter(col)][1:]:
            c.number_format = 'DD/MM/YYYY'
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical='top')


ws = wb.active
ws.title = 'Itens revogados'
sheet(ws, ['Anexo', 'Segmento (Anexo)', 'Item', 'CEST', 'NCM/SH', 'Descrição', 'Tipo de revogação', 'Ato revogador',
           'Data do ato', 'Publicação DOE', 'Data de revogação (início da vigência)', 'Situação em 30/09/2026'],
      [[d['anexo'], d['anexo_titulo'], d['item'], d['cest'], d['ncm'], d['descricao'], d['tipo'], 'Portaria ' + d['ato'],
        d['data_ato'], d['doe'], d['vig'], situacao(d)] for d in itens],
      [8, 34, 7, 11, 16, 70, 22, 16, 12, 14, 18, 30], datecols=(9, 10, 11))

cnt = Counter(d['anexo'] for d in itens)
ws2 = wb.create_sheet('Anexos revogados')
data = []
for ax, (t, ato) in sorted(annexes.items(), key=lambda kv: ORDEM[kv[0]]):
    a = ACTS[ato]
    data.append([ax, t, 'Portaria ' + ato, a[0], a[1], a[2], cnt.get(ax, 0)])
sheet(ws2, ['Anexo', 'Segmento', 'Ato revogador', 'Data do ato', 'Publicação DOE', 'Data de revogação',
            'Itens do anexo na planilha (inclui os revogados antes, individualmente)'],
      data, [8, 50, 16, 12, 14, 16, 30], datecols=(4, 5, 6))


def arquivo(k):
    m = re.match(r'^(\w+)-(\d+)/(\d+)$', k)
    return 'atos-revogadores/Portaria-%s-%d-de-20%s.html' % (m.group(1), int(m.group(2)), m.group(3))


ws3 = wb.create_sheet('Atos revogadores')
sheet(ws3, ['Ato', 'Data do ato', 'Publicação DOE', 'Início da vigência', 'Fonte da data de vigência', 'Arquivo local'],
      [['Portaria ' + k, v[0], v[1], v[2], v[3], arquivo(k)] for k, v in sorted(ACTS.items(), key=lambda kv: kv[1][2])],
      [16, 12, 14, 16, 55, 50], datecols=(2, 3, 4))

ws4 = wb.create_sheet('Leia-me')
for line in [
    'Produtos revogados da Portaria CAT 68/2019 (SP) — relação de mercadorias sujeitas à ST com retenção antecipada.',
    'Uso: base para o levantamento de crédito de ICMS sobre estoque na exclusão da ST (Portaria CAT 28/2020).',
    'Fonte: https://legislacao.fazenda.sp.gov.br/Paginas/Portaria-CAT-68-de-2019.aspx (texto consolidado; cópia em portaria-cat-68-2019.html).',
    'Extração em 30/09/2026. A página consolidada exibia data 06/07/2026 — conferir atos posteriores.',
    '"Data de revogação" = início da vigência da revogação (exclusão da ST). Estoque a inventariar (CAT 28/2020, art. 2º) = o do final do dia anterior a essa data.',
    'Tipo "Item revogado": item revogado individualmente (vale a data do ato que o revogou).',
    'Tipo "Anexo inteiro revogado": item sem revogação individual em anexo revogado integralmente (vale a data de revogação do anexo).',
    'Item revogado individualmente antes da revogação do anexo mantém a data individual (mais antiga).',
    'Descrição/NCM/CEST = redação corrente do item no consolidado (1ª linha do item; as demais são redações anteriores).',
    'Revogações com data posterior a 30/09/2026 estão marcadas como "futura".',
    'Anexo XI: o consolidado cita "DOE 23-10-2025" para a Portaria SRE-94/25; o ato oficial informa publicação em 23/12/2025 (usada aqui).',
    'Antes de usar no levantamento: validar NCM/CEST e se o produto foi realocado em outro anexo/regime, além de regras de transição.',
]:
    ws4.append([line])
ws4.column_dimensions['A'].width = 150
ws4['A1'].font = Font(bold=True, size=12)

wb.save(os.path.join(BASE, 'cat68-2019-produtos-revogados.xlsx'))

with open(os.path.join(BASE, 'cat68-2019-produtos-revogados.csv'), 'w', newline='', encoding='utf-8-sig') as f:
    w = csv.writer(f, delimiter=';')
    w.writerow(['anexo', 'segmento', 'item', 'cest', 'ncm', 'descricao', 'tipo_revogacao', 'ato_revogador', 'data_ato',
                'doe', 'data_revogacao', 'situacao'])
    for d in itens:
        w.writerow([d['anexo'], d['anexo_titulo'], d['item'], d['cest'], d['ncm'], d['descricao'], d['tipo'],
                    'Portaria ' + d['ato'], d['data_ato'].strftime('%d/%m/%Y'), d['doe'].strftime('%d/%m/%Y'),
                    d['vig'].strftime('%d/%m/%Y'), situacao(d)])

# Itens que continuam na ST (sem revogação própria nem do anexo): serve para detectar CEST/NCM que ainda estão na ST
rev_keys = {(d['anexo'], d['item']) for d in itens}
vigentes = [r for k, r in cur.items() if k not in rev_keys]
vigentes.sort(key=lambda r: (ORDEM[r['anexo']], sortnum(r['item'])))
with open(os.path.join(BASE, 'cat68-2019-itens-vigentes.csv'), 'w', newline='', encoding='utf-8-sig') as f:
    w = csv.writer(f, delimiter=';')
    w.writerow(['anexo', 'segmento', 'item', 'cest', 'ncm', 'descricao'])
    for r in vigentes:
        w.writerow([r['anexo'], r['anexo_titulo'], r['item'], r['cest'], r['ncm'], r['descricao']])

print(len(itens), 'itens revogados')
print(len(vigentes), 'itens vigentes na ST')
print(sorted(Counter((d['anexo'], d['tipo']) for d in itens).items(), key=lambda kv: ORDEM[kv[0][0]]))
print(sorted(Counter((d['ato'], d['vig'].isoformat()) for d in itens).items()))
print('anexos revogados:', list(annexes))
