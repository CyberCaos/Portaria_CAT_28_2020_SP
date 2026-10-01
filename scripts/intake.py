"""Etapa 1 do levantamento: lê o pacote do usuário (pasta ou .zip com XMLs de compras + posições de estoque) e
consolida tudo em `entrada_consolidada.xlsx` (+ CSVs e resumo JSON). Não calcula crédito: só prepara e audita a entrada.

Uso:
    python scripts/intake.py --entrada <pasta|arquivo.zip> --saida <pasta_de_saida> [--cnpj 00000000000000]

Estrutura aceita (qualquer subpasta, zips aninhados também):
    compras/   *.xml  (NF-e/NFC-e de entrada, procEventoNFe de cancelamento)
    estoque/   estoque 31.12.2025.xlsx, estoque 31.03.2026.xlsx ...  (data no nome do arquivo)
"""
import argparse
import hashlib
import json
import os
import re
import sys
from collections import Counter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import compras as C  # noqa: E402
import estoque as E  # noqa: E402
import triagem as PE  # noqa: E402   (resolver_cest / indice_cest_por_ean)
from parser_xml.factory import listar_xml, preparar_origem, processar_lote  # noqa: E402


def _digitos(s):
    return re.sub(r'\D', '', s or '')


def consolidar(origem, saida: str, cnpj: str = '', progresso=None) -> dict:
    """origem: pasta/zip ou lista de pastas/zips (XMLs e posições de estoque em qualquer um deles)."""
    os.makedirs(saida, exist_ok=True)
    origens = [origem] if isinstance(origem, (str, os.PathLike)) else list(origem)   # várias pastas/zips são aceitos
    raizes, limpezas = [], []
    for o in origens:
        r, l = preparar_origem(str(o))
        raizes += r
        limpezas.append(l)
    limpar = lambda: [l() for l in limpezas]
    try:
        xmls = [c for r in raizes for c in listar_xml(r)]
        lote = processar_lote(xmls, progresso)
        arquivos_estoque, vistos = [], set()
        for c, d in E.descobrir(raizes):      # o mesmo arquivo pode aparecer em mais de uma origem: lê uma vez só
            with open(c, 'rb') as fh:
                chave = (d, os.path.basename(c).lower(), hashlib.sha1(fh.read()).hexdigest())
            if chave not in vistos:
                vistos.add(chave)
                arquivos_estoque.append((c, d))
        leituras = [E.ler(c, d) for c, d in arquivos_estoque]
    finally:
        limpar()

    cnpj = _digitos(cnpj)
    linhas = []
    for it in lote.itens:
        r = C.linha_compra(it)
        if cnpj:
            r['papel_empresa'] = ('entrada' if _digitos(it.destinatario_documento) == cnpj else
                                  'saida' if _digitos(it.emitente_documento) == cnpj else 'terceiros')
        linhas.append(r)
    campos = C.CAMPOS_COMPRAS + (['papel_empresa'] if cnpj else [])
    linhas.sort(key=lambda r: (r['data_emissao'] or __import__('datetime').date.min, r['chave_nfe'], r['n_item']))
    estoque_linhas = [l for lei in leituras for l in lei['linhas']]
    rejeitadas = [r for lei in leituras for r in lei['rejeitadas']]
    cobertura_cest = PE.resolver_cest(estoque_linhas, PE.indice_cest_por_ean(linhas))
    # posições com datas diferentes e conteúdo idêntico = exportação não histórica (todas refletem o mesmo momento)
    grupos = {}
    for lei in leituras:
        if lei['linhas']:
            grupos.setdefault(E.assinatura(lei['linhas']), []).append(lei)
    identicas = [[(g['arquivo'], str(g['data_posicao'])) for g in gs] for gs in grupos.values()
                 if len({g['data_posicao'] for g in gs}) > 1]

    notas = {r['chave_nfe'] for r in linhas}
    com_st_ret = [r for r in linhas if r['vbc_st_ret'] > 0]
    cst_st = [r for r in linhas if r['cst'] in ('60', '500', '10', '70', '30', '201', '202', '203', '900')]
    datas = [r['data_emissao'] for r in linhas if r['data_emissao']]
    resumo = {
        'entrada': [str(o) for o in origens],
        'xml_arquivos_lidos': lote.arquivos,
        'xml_com_erro': len(lote.erros),
        'xml_ignorados_aviso': len(lote.avisos),
        'notas_unicas': len(notas),
        'notas_duplicadas_descartadas': lote.notas_duplicadas,
        'notas_canceladas': lote.notas_canceladas,
        'eventos_cancelamento': len(lote.eventos),
        'eventos_orfaos': len(lote.eventos_orfaos),
        'itens_de_compra': len(linhas),
        'itens_com_bc_st_retida': len(com_st_ret),
        'itens_st_sem_bc_st_retida': sum(1 for r in cst_st if r['cst'] in ('60', '500') and r['vbc_st_ret'] <= 0),
        'notas_sem_protocolo': len({r['chave_nfe'] for r in linhas if not r['c_stat']}),
        'notas_nao_autorizadas': len({r['chave_nfe'] for r in linhas if r['c_stat'] and r['c_stat'] not in ('100', '150')}),
        'emissao_primeira': str(min(datas)) if datas else None,
        'emissao_ultima': str(max(datas)) if datas else None,
        'destinatarios_distintos': len({r['cnpj_destinatario'] for r in linhas}),
        'papel_empresa': dict(Counter(r.get('papel_empresa') for r in linhas)) if cnpj else None,
        'arquivos_estoque': [
            {'arquivo': l['arquivo'], 'data_posicao': str(l['data_posicao']) if l['data_posicao'] else None,
             'linhas': len(l['linhas']), 'rejeitadas': len(l['rejeitadas']), 'total_rodape': l['total_rodape'],
             'colunas_detectadas': l['colunas'], 'nao_reconhecidas': l['nao_reconhecidas'],
             'bloqueantes': l['bloqueantes'], 'avisos': l['avisos']} for l in leituras],
        'estoque_linhas_total': len(estoque_linhas),
        'estoque_linhas_rejeitadas': len(rejeitadas),
        'estoque_cobertura_cest': cobertura_cest,
        'estoque_posicoes_identicas': identicas,
    }
    alertas = []
    if not xmls:
        alertas.append('Nenhum XML encontrado.')
    if not leituras:
        alertas.append('Nenhum arquivo de estoque encontrado (nome precisa conter "estoque").')
    if not cnpj:
        alertas.append('CNPJ da empresa não informado: não foi possível separar entradas de saídas.')
    elif resumo['destinatarios_distintos'] > 1 and resumo['papel_empresa'].get('entrada', 0) == 0:
        alertas.append('Nenhuma nota tem o CNPJ informado como destinatário: conferir CNPJ/pacote.')
    for l in leituras:
        for b in l['bloqueantes']:
            alertas.append(f"BLOQUEANTE estoque '{l['arquivo']}': {b}")
    if rejeitadas:
        alertas.append(f'{len(rejeitadas)} linha(s) de estoque rejeitada(s) por EAN/Descrição/NCM/Valor ausente ou '
                       'inválido (aba "Estoque rejeitadas"): corrigir na origem ou aceitar a exclusão da análise.')
    if identicas:                                    # informativo: a posição enviada vale como está (decisão do usuário)
        alertas.append('Info: posições de datas diferentes com o mesmo conteúdo ('
                       + '; '.join(' = '.join(f'{a} ({d})' for a, d in g) for g in identicas)
                       + '); usadas como enviadas pelo cliente.')
    if estoque_linhas and cobertura_cest['total']:
        sem = cobertura_cest['nao_encontrado'] + cobertura_cest['nfe_ean_ambiguo']
        if sem:
            alertas.append(f"CEST não resolvido com segurança para {sem} de {cobertura_cest['total']} linhas de estoque "
                           f"({cobertura_cest['nao_encontrado']} sem nota com esse EAN, "
                           f"{cobertura_cest['nfe_ean_ambiguo']} com CEST ambíguo): ficam em revisão no cruzamento "
                           'com a CAT 68.')
    if lote.erros:
        alertas.append(f'{len(lote.erros)} XML com erro (ver aba "Erros XML").')
    if resumo['itens_st_sem_bc_st_retida']:
        alertas.append(f"{resumo['itens_st_sem_bc_st_retida']} itens CST 60/CSOSN 500 sem vBCSTRet: crédito zero até NF complementar (art. 4º).")
    if lote.eventos_orfaos:
        alertas.append(f'{len(lote.eventos_orfaos)} evento(s) de cancelamento sem a nota no pacote.')
    resumo['alertas'] = alertas

    _escrever(saida, resumo, campos, linhas, estoque_linhas, rejeitadas, lote)
    return resumo


def _escrever(saida, resumo, campos, linhas, estoque_linhas, rejeitadas, lote):
    import csv
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    hdr = PatternFill('solid', fgColor='1F3864')

    def aba(ws, cab, dados, larguras=None, fmt_data=()):
        ws.append(cab)
        for c in ws[1]:
            c.fill, c.font, c.alignment = hdr, Font(bold=True, color='FFFFFF'), Alignment(wrap_text=True)
        for d in dados:
            ws.append(d)
        for j, nome in enumerate(cab, 1):
            ws.column_dimensions[get_column_letter(j)].width = (larguras or {}).get(nome, max(10, min(len(nome) + 3, 40)))
        for j in fmt_data:
            for c in ws[get_column_letter(j)][1:]:
                c.number_format = 'DD/MM/YYYY'
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions

    ws = wb.active
    ws.title = 'Resumo'
    ws.append(['Indicador', 'Valor'])
    for c in ws[1]:
        c.fill, c.font = hdr, Font(bold=True, color='FFFFFF')
    for k, v in resumo.items():
        if k in ('arquivos_estoque', 'alertas'):
            continue
        ws.append([k, json.dumps(v, ensure_ascii=False) if isinstance(v, (dict, list)) else v])
    ws.append([])
    ws.append(['ALERTAS'])
    ws[ws.max_row][0].font = Font(bold=True)
    for a in resumo['alertas'] or ['(nenhum)']:
        ws.append([a])
    ws.append([])
    ws.append(['ARQUIVOS DE ESTOQUE'])
    ws[ws.max_row][0].font = Font(bold=True)
    for a in resumo['arquivos_estoque']:
        ws.append([a['arquivo'], f"data={a['data_posicao']} linhas={a['linhas']} rejeitadas={a['rejeitadas']} rodape={a['total_rodape']}",
                   'colunas: ' + json.dumps(a['colunas_detectadas'], ensure_ascii=False),
                   'não reconhecidas: ' + ', '.join(a['nao_reconhecidas']), ' | '.join(a['avisos'])])
    ws.column_dimensions['A'].width = 42
    ws.column_dimensions['B'].width = 40

    aba(wb.create_sheet('Compras'), campos, [[r.get(c) for c in campos] for r in linhas],
        {'descricao': 40, 'nome_emitente': 30, 'chave_nfe': 46}, fmt_data=(campos.index('data_emissao') + 1,))
    ecampos = ['data_posicao', 'arquivo', 'ean', 'cod_produto', 'descricao', 'ncm', 'cest', 'cest_origem', 'qtd',
               'unidade', 'valor', 'valor_total', 'grupo_pai', 'cst_icms', 'totalizador', 'aliq_totalizador', 'ean_dv_ok']
    aba(wb.create_sheet('Estoque'), ecampos, [[r.get(c) for c in ecampos] for r in estoque_linhas],
        {'descricao': 40, 'arquivo': 28, 'ean': 16}, fmt_data=(1,))
    aba(wb.create_sheet('Estoque rejeitadas'), ['arquivo', 'linha', 'ean', 'descricao', 'motivo'],
        [[r['arquivo'], r['linha'], r['ean'], r['descricao'], r['motivo']] for r in rejeitadas],
        {'arquivo': 28, 'descricao': 40, 'motivo': 70})
    aba(wb.create_sheet('Erros XML'), ['arquivo', 'erro'], [[e.nome, e.erro] for e in lote.erros], {'arquivo': 40, 'erro': 90})
    aba(wb.create_sheet('Avisos XML'), ['arquivo', 'aviso'], [[e.nome, e.aviso] for e in lote.avisos], {'arquivo': 40, 'aviso': 90})
    aba(wb.create_sheet('Eventos'), ['chave_nfe', 'data_evento', 'arquivo', 'orfao'],
        [[e.ch_nfe, e.dh_evento, e.nome_arquivo, 'sim' if e in lote.eventos_orfaos else 'não'] for e in lote.eventos],
        {'chave_nfe': 46, 'arquivo': 40})
    wb.save(os.path.join(saida, 'entrada_consolidada.xlsx'))

    for nome, cab, dados in (('compras.csv', campos, [[r.get(c) for c in campos] for r in linhas]),
                             ('estoque.csv', ecampos, [[r.get(c) for c in ecampos] for r in estoque_linhas])):
        with open(os.path.join(saida, nome), 'w', newline='', encoding='utf-8-sig') as f:
            w = csv.writer(f, delimiter=';')
            w.writerow(cab)
            w.writerows(dados)
    with open(os.path.join(saida, 'intake_resumo.json'), 'w', encoding='utf-8') as f:
        json.dump(resumo, f, ensure_ascii=False, indent=2, default=str)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--entrada', required=True, nargs='+', help='uma ou mais pastas/.zip com XMLs e posições de estoque')
    ap.add_argument('--saida', required=True)
    ap.add_argument('--cnpj', default='', help='CNPJ da empresa analisada (separa entradas de saídas)')
    a = ap.parse_args()
    r = consolidar(a.entrada if len(a.entrada) > 1 else a.entrada[0], a.saida, a.cnpj, lambda i, n: print(f'  {i}/{n} XML...', flush=True))
    print(json.dumps({k: v for k, v in r.items() if k != 'arquivos_estoque'}, ensure_ascii=False, indent=2, default=str))
    for l in r['arquivos_estoque']:
        print('estoque:', l['arquivo'], l['data_posicao'], l['linhas'], 'linhas', l['avisos'] or '')


if __name__ == '__main__':
    main()
