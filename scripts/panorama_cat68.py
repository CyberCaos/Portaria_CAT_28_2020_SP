"""Panorama das saídas da ST (Portaria CAT 68/2019) em um ano: o que já saiu, o que sai e quais posições de estoque levantar.

Lê `references/base-legal/cat68-2019-produtos-revogados.csv` (itens revogados, com a data de efeito) e
`cat68-2019-itens-vigentes.csv` (itens que continuam na ST) e grava uma planilha:

    python scripts/panorama_cat68.py [--ano 2026] [--saida <pasta>] [--hoje AAAA-MM-DD]

Abas: Resumo por data, Por anexo, Itens (um por linha, com NCM, CEST e critério de triagem), Posições de estoque
(o que pedir ao cliente) e Ainda na ST. A situação (já saiu / vai sair) é calculada pela data de hoje, não pelo rótulo
da base, que foi gravado na data da extração.
"""
import argparse
import csv
import os
import sys
from collections import defaultdict
from datetime import date, datetime, timedelta

BASE = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'references', 'base-legal')


def _csv(nome):
    with open(os.path.join(BASE, nome), encoding='utf-8-sig', newline='') as f:
        return [{(k or '').lstrip('﻿'): v for k, v in r.items()} for r in csv.DictReader(f, delimiter=';')]


def _data(s):
    return datetime.strptime(s.strip(), '%d/%m/%Y').date()


def _situacao(d, hoje):
    dias = (d - hoje).days
    if dias < 0:
        return f'Já saiu (há {-dias} dias)'
    if dias == 0:
        return 'Sai hoje'
    return f'Vai sair (em {dias} dias)'


def carregar(ano):
    rev = [r for r in _csv('cat68-2019-produtos-revogados.csv') if r['data_revogacao'].endswith(str(ano))]
    completos = {r['anexo'] for r in _csv('cat68-2019-produtos-revogados.csv') if r['tipo_revogacao'] == 'Anexo inteiro revogado'}
    for r in rev:
        r['_data'] = _data(r['data_revogacao'])
        r['_criterio'] = 'só NCM' if r['anexo'] in completos else 'NCM + CEST'
    return rev, _csv('cat68-2019-itens-vigentes.csv')


def montar(ano, hoje):
    rev, vig = carregar(ano)
    if not rev:
        raise ValueError(f'Nenhuma revogação com data em {ano} na base.')
    datas = sorted({r['_data'] for r in rev})
    resumo, posicoes = [], []
    for d in datas:
        itens = [r for r in rev if r['_data'] == d]
        anexos = sorted({r['anexo'] for r in itens}, key=lambda a: (len(a), a))
        atos = sorted({r['ato_revogador'] for r in itens})
        criterios = sorted({r['_criterio'] for r in itens})
        pos = d - timedelta(days=1)
        resumo.append([d.strftime('%d/%m/%Y'), _situacao(d, hoje), '; '.join(atos), ', '.join(anexos), len(itens),
                       ' / '.join(criterios), pos.strftime('%d/%m/%Y'),
                       '; '.join(sorted({r['segmento'].title() for r in itens}))])
        posicoes.append([pos.strftime('%d/%m/%Y'), f"estoque {pos.strftime('%d.%m.%Y')}.xls", d.strftime('%d/%m/%Y'), ', '.join(anexos), len(itens),
                         'Já era devida' if pos < hoje else ('Fechar o estoque neste dia' if pos > hoje else 'Fechar hoje')])
    por_anexo = defaultdict(lambda: dict(seg='', itens=0, datas=set(), crit=set()))
    for r in rev:
        a = por_anexo[r['anexo']]
        a['seg'] = r['segmento'].title()
        a['itens'] += 1
        a['datas'].add(r['_data'])
        a['crit'].add(r['_criterio'])
    anexos = [[k, v['seg'], v['itens'], ', '.join(d.strftime('%d/%m/%Y') for d in sorted(v['datas'])), ' / '.join(sorted(v['crit']))]
              for k, v in sorted(por_anexo.items(), key=lambda kv: (len(kv[0]), kv[0]))]
    itens = [[r['_data'].strftime('%d/%m/%Y'), _situacao(r['_data'], hoje), r['anexo'], r['segmento'].title(), r['item'], r['cest'], r['ncm'],
              r['descricao'], r['_criterio'], r['tipo_revogacao'], r['ato_revogador'], r['doe']]
             for r in sorted(rev, key=lambda r: (r['_data'], len(r['anexo']), r['anexo'], r['item']))]
    na_st = [[r['anexo'], r['segmento'].title(), r['item'], r['cest'], r['ncm'], r['descricao']] for r in vig]
    return dict(rev=rev, vig=vig, resumo=resumo, anexos=anexos, itens=itens, posicoes=posicoes, na_st=na_st, datas=datas)


def gravar(ano, hoje, pasta):
    import pandas as pd
    from openpyxl.styles import Alignment, Font, PatternFill
    p = montar(ano, hoje)
    abas = {
        'Resumo por data': (['Data da saída da ST', 'Situação', 'Ato revogador', 'Anexos', 'Itens', 'Critério de triagem',
                              'Posição de estoque necessária', 'Segmentos'], p['resumo']),
        'Por anexo': (['Anexo', 'Segmento', 'Itens', 'Data(s) da saída', 'Critério de triagem'], p['anexos']),
        'Itens': (['Data da saída', 'Situação', 'Anexo', 'Segmento', 'Item', 'CEST', 'NCM', 'Descrição', 'Critério', 'Tipo da revogação',
                   'Ato revogador', 'DOE'], p['itens']),
        'Posições de estoque': (['Posição (fim do dia)', 'Arquivo esperado', 'Atende a saída de', 'Anexos', 'Itens', 'Situação'], p['posicoes']),
        'Ainda na ST': (['Anexo', 'Segmento', 'Item', 'CEST', 'NCM', 'Descrição'], p['na_st']),
    }
    os.makedirs(pasta, exist_ok=True)
    destino = os.path.join(pasta, f'panorama_cat68_{ano}.xlsx')
    with pd.ExcelWriter(destino, engine='openpyxl') as w:
        for nome, (cols, linhas) in abas.items():
            pd.DataFrame(linhas, columns=cols).to_excel(w, sheet_name=nome, index=False)
            ws = w.sheets[nome]
            ws.freeze_panes = 'A2'
            for c in ws[1]:
                c.font = Font(bold=True, color='FFFFFF')
                c.fill = PatternFill('solid', fgColor='1F3A5F')
                c.alignment = Alignment(wrap_text=True, vertical='center')
            for col in ws.columns:
                tam = max(len(str(c.value or '')) for c in col[:200])
                ws.column_dimensions[col[0].column_letter].width = min(max(12, tam + 2), 70)
            if nome in ('Itens', 'Ainda na ST'):
                ws.auto_filter.ref = ws.dimensions
    return destino, p


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--ano', type=int, default=2026)
    ap.add_argument('--saida', default='.', help='pasta onde gravar a planilha')
    ap.add_argument('--hoje', default='', help='data de referência AAAA-MM-DD (padrão: hoje)')
    a = ap.parse_args(argv)
    hoje = datetime.strptime(a.hoje, '%Y-%m-%d').date() if a.hoje else date.today()
    try:
        destino, p = gravar(a.ano, hoje, a.saida)
    except ValueError as exc:
        sys.exit(f'Erro: {exc}')
    print(f'{destino}\n')
    print(f"Saídas da ST em {a.ano} (referência {hoje.strftime('%d/%m/%Y')}): {len(p['rev'])} itens em {len(p['datas'])} datas")
    for r in p['resumo']:
        print(f"  {r[0]}  {r[1]:<26} {r[4]:>4} itens  anexos {r[3]}  (posição de {r[6]})")
    futuras = [d for d in p['datas'] if d > hoje]
    print(f"\nAinda na ST: {len(p['vig'])} itens, sem data de saída registrada na base.")
    if not futuras:
        print('Nenhuma saída futura registrada na base: confira atos novos antes de concluir (references/base-legal/README.md).')


if __name__ == '__main__':
    main()
