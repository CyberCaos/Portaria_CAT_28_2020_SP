"""Levantamento — organização dos dados (passos 1 a 5 do usuário) sobre a pasta do cliente.

    python scripts/levantamento.py <pasta_do_cliente>

Entrada : <cliente>/relatorio_parser/compras.csv e estoque.csv (gerados por preparar_cliente.py / intake.py)
          <cliente>/cliente.json (regime, nome, CNPJ)
Saída   : <cliente>/resultado_levantamento.xlsx (planilha do cliente): Resumo (com premissas), Anexo II (notas
          selecionadas), Resumo por produto, Resumo por competência, Pendências.
          <cliente>/relatorio_parser/trilha_levantamento.xlsx (trilha técnica, lida pelo PDF): Triagem estoque, Localização,
          Alocação de notas, Fora do crédito, Uso de notas, Premissas.
Travas: só nota provada gera crédito (sem validação manual) e cada item de nota supre no máximo a quantidade comprada,
somando todas as posições (livro-razão único; itens casados por EAN escolhem primeiro).
"""
import csv
import json
import os
import sys
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import alocacao as A  # noqa: E402
import credito as CR  # noqa: E402
import enquadramento_auto as EA  # noqa: E402
import identidade_cat68 as ID  # noqa: E402
import fator_conversao as F  # noqa: E402
import localizar as L  # noqa: E402
import resultado_modelo as RM  # noqa: E402
import triagem as T  # noqa: E402

NUM_COMPRAS = ('qtd', 'qtd_trib', 'vl_merc', 'frete', 'seguro', 'outras_despesas', 'desconto', 'vl_contabil', 'vbc_st_ret', 'p_st',
               'vicms_st_ret', 'vbc_fcp_st_ret', 'vicms_substituto', 'vbc_st', 'vicms_st', 'p_icms_st', 'p_fcp_st', 'p_fcp_st_ret',
               'p_icms', 'p_red_bc', 'p_red_bc_st', 'bc_st_antecipacao', 'vbc', 'vicms', 'vfcp_st', 'vfcp_st_ret', 'vbc_fcp_st')


def _ler_csv(caminho):
    with open(caminho, encoding='utf-8-sig', newline='') as f:
        return list(csv.DictReader(f, delimiter=';'))


def _num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def _data(s):
    return datetime.strptime(s[:10], '%Y-%m-%d').date()


def _fmt(d):
    return d.strftime('%d/%m/%Y') if isinstance(d, date) else d


def carregar(pasta: str):
    rel = os.path.join(pasta, 'relatorio_parser')
    compras = _ler_csv(os.path.join(rel, 'compras.csv'))
    campos = list(compras[0].keys()) if compras else []
    faltam = [c for c in OBRIGATORIAS_COMPRAS if compras and c not in campos]
    if faltam:
        raise ValueError(f'compras.csv sem as colunas fiscais {faltam}: gerado por versão antiga do parser. '
                         'Rode de novo python scripts/preparar_cliente.py com o mesmo pacote (o cálculo sairia errado sem elas).')
    for c in compras:
        for k in NUM_COMPRAS:
            if k in c:
                c[k] = _num(c[k])
    estoque = _ler_csv(os.path.join(rel, 'estoque.csv'))
    for e in estoque:
        e['qtd'], e['valor'], e['valor_total'] = _num(e['qtd']), _num(e['valor']), _num(e['valor_total'])
        e['aliq_totalizador'] = _num(e['aliq_totalizador']) if e.get('aliq_totalizador') not in (None, '') else None
    with open(os.path.join(pasta, 'cliente.json'), encoding='utf-8') as f:
        cli = json.load(f)
    with open(os.path.join(rel, 'intake_resumo.json'), encoding='utf-8') as f:
        resumo = json.load(f)
    return compras, campos, estoque, cli, resumo


OBRIGATORIAS_COMPRAS = ('cst', 'orig', 'crt_emitente', 'uf_emitente', 'uf_destinatario', 'vl_merc', 'frete', 'seguro', 'outras_despesas',
                        'desconto', 'vbc_st', 'vbc_st_ret', 'vicms_st', 'vicms_st_ret', 'vicms_substituto', 'p_icms_st', 'p_fcp_st', 'p_st',
                        'p_icms', 'vicms', 'vfcp_st', 'vfcp_st_ret', 'vbc_fcp_st', 'vbc_fcp_st_ret', 'ref_nfe', 'fin_nfe', 'c_stat',
                        'p_red_bc', 'p_red_bc_st')
ENQUADRAMENTO = 'enquadramento_reducao.csv'
CAMPOS_ENQ = ['tipo', 'codigo', 'descricao', 'ncm', 'cest', 'p_red_bc', 'origem_reducao', 'reducao', 'dispositivo', 'evidencia_tipo',
              'evidencia_fonte', 'apresentacao_confirmada', 'vigencia_condicoes', 'justificativa', 'fonte', 'sugestao']
# evidência que sustenta a identificação do produto (não importa QUEM analisou, e sim O QUE prova)
EVIDENCIAS_VALIDAS = ('descricao_principio_ativo', 'gtin_registro_sanitario', 'documentacao_fabricante', 'documento_oficial',
                      'descricao_e_classificacao')
ENTRADAS = 'entradas.csv'
IDENTIDADE = 'identidade_cat68.csv'
CAMPOS_ID = ['ean', 'descricao', 'ncm', 'cest', 'anexo', 'item', 'descricao_legal', 'motivo', 'confirmado', 'evidencia_tipo',
             'evidencia_fonte', 'apresentacao_confirmada', 'justificativa', 'fonte', 'sugestao']


def _evidencia_ok(r, exige_dispositivo=True):
    """Evidência completa de uma linha preenchida (enquadramento ou identidade). Devolve (ok, o que falta)."""
    falta = []
    if exige_dispositivo and not str(r.get('dispositivo') or '').strip():
        falta.append('dispositivo')
    if str(r.get('evidencia_tipo') or '').strip().lower() not in EVIDENCIAS_VALIDAS:
        falta.append('tipo de evidência (' + ', '.join(EVIDENCIAS_VALIDAS) + ')')
    if not str(r.get('evidencia_fonte') or '').strip():
        falta.append('fonte da evidência (registro, link, documento)')
    if str(r.get('apresentacao_confirmada') or '').strip().lower() not in ('sim', 's'):
        falta.append('confirmação de que a evidência é da apresentação exata (EAN/descrição)')
    if exige_dispositivo and str(r.get('vigencia_condicoes') or '').strip().lower() not in ('sim', 's'):
        falta.append('vigência e condições do dispositivo verificadas')
    return not falta, '; '.join(falta)
VALORES_ENQ = {'aplicavel': 'reducao_aplicavel_consumidor_final', 'nao_aplicavel': 'reducao_nao_aplicavel_consumidor_final'}
ROTULO_ENQ = {v: k for k, v in VALORES_ENQ.items()}
FONTE_AUTO = 'automatico'


def _brl(v):
    return 'R$ ' + f'{v:,.2f}'.replace(',', 'X').replace('.', ',').replace('X', '.')


def _dig(v):
    return ''.join(ch for ch in str(v or '') if ch.isdigit())


def _tem_reducao(c):
    return float(c.get('p_red_bc_st') or 0) > 0 or (float(c.get('p_red_bc') or 0) > 0 and float(c.get('vbc_st') or 0) > 0)


def _carga_do_dispositivo(texto):
    """Carga legal (%) do dispositivo citado numa regra do usuário/análise, pela tabela do RICMS/SP (None se não fixa carga)."""
    t = str(texto or '').strip().lower()
    if not t:
        return None
    for regra in EA.tabela():
        if regra['_carga'] is not None and regra['dispositivo'].lower() in t:
            return regra['_carga']
    return None


def _chave_enq(c):
    ean = L.gtin14(c.get('ean'))
    return ('EAN', _dig(ean)) if ean else ('NCM', _dig(c.get('ncm')))


def _carregar_enquadramento(pasta):
    """<cliente>/enquadramento_reducao.csv: tipo (EAN, CEST ou NCM), codigo, reducao (aplicavel | nao_aplicavel), dispositivo,
    justificativa e fonte. Devolve (linhas a preservar, regras do usuário/análise por tipo, decisões automáticas anteriores):
    - preenchidas com fonte diferente de "automatico": regras do usuário ou da análise, prevalecem;
    - fonte "automatico": refeitas a cada rodada; se o usuário trocou a reducao sem trocar a fonte, a troca vale como regra dele;
    - não automáticas com dispositivo ou justificativa mas sem reducao: análise em andamento, preservadas (sem virar regra)."""
    caminho = os.path.join(pasta, ENQUADRAMENTO)
    if not os.path.exists(caminho):
        return [], {}, {}
    linhas, auto_ant = [], {}
    regras = defaultdict(dict)
    for r in _ler_csv(caminho):
        fonte = str(r.get('fonte') or '').strip().lower()
        valor = VALORES_ENQ.get(str(r.get('reducao') or '').strip().lower())
        tipo, cod = str(r.get('tipo') or '').strip().upper(), _dig(r.get('codigo'))
        if fonte == FONTE_AUTO:
            if tipo and cod:
                auto_ant[(tipo, cod)] = r
            continue
        if valor and tipo in ('EAN', 'CEST', 'NCM') and cod:
            fonte_r = 'analise' if fonte == 'analise' else 'usuario'
            ok, falta = _evidencia_ok(r)
            regras[tipo][L.gtin14(cod) if tipo == 'EAN' else cod] = (valor, _carga_do_dispositivo(r.get('dispositivo')), fonte_r,
                                                                     str(r.get('dispositivo') or ''), ok, falta)
            linhas.append(r)
        elif str(r.get('dispositivo') or '').strip() or str(r.get('justificativa') or '').strip():
            linhas.append(r)                                         # análise em andamento: não apagar
    return linhas, regras, auto_ant


def _aplicar_enquadramento(compras, regras, auto_ant=None, regime=''):
    """Grava c['reducao'] (e a carga legal, quando o dispositivo a fixa) nas linhas de compra COM redução: primeiro a regra do
    usuário/análise (EAN da nota > CEST > NCM, o NCM do arquivo pode ter menos dígitos); sem regra, o automático pelo RICMS/SP
    (enquadramento_auto.py). Linha automática que o usuário editou sem trocar a fonte vale como regra dele."""
    auto_ant = auto_ant or {}
    n = 0
    for c in compras:
        if not _tem_reducao(c):                                      # nota sem redução do mesmo produto segue sem redução
            continue
        achado = regras.get('EAN', {}).get(L.gtin14(c.get('ean')) or '-') or regras.get('CEST', {}).get(_dig(c.get('cest')) or '-')
        if not achado:
            ncm = _dig(c.get('ncm'))
            for k in sorted(regras.get('NCM', {}), key=len, reverse=True):
                if ncm.startswith(k):
                    achado = regras['NCM'][k]
                    break
        valor, carga, fonte_enq, disp, evid_ok, evid_falta = achado if achado else (None, None, '', '', False, '')
        if not valor:
            auto = EA.classificar(c, regime)
            ant = auto_ant.get(_chave_enq(c))
            editado = VALORES_ENQ.get(str((ant or {}).get('reducao') or '').strip().lower())
            if ant and editado and auto and editado != auto['reducao']:      # o automático decide X e o arquivo diz Y: edição do usuário
                valor, carga = editado, _carga_do_dispositivo(ant.get('dispositivo'))
                fonte_enq, disp = 'usuario', str(ant.get('dispositivo') or '')
                evid_ok, evid_falta = _evidencia_ok(ant)
                c['_enq_editado'] = ant                              # o usuário mudou a decisão automática
            elif auto:
                valor, carga = auto['reducao'], auto.get('carga')
                fonte_enq, disp = 'automatico', auto['dispositivo']
                evid_ok, evid_falta = True, ''                      # princípio ativo escrito, NCM, UF e vigência conferidos
                c['_enq_auto'] = auto
                c['_enq_divergente'] = auto.get('divergente', False)
        if valor:
            c['reducao'] = valor
            c['_enq_fonte'], c['_enq_dispositivo'] = fonte_enq, disp
            c['_enq_evid_ok'], c['_enq_evid_falta'] = evid_ok, evid_falta
            if carga is not None:
                c['_carga_legal'] = carga
            n += 1
    return n


def _gravar_enquadramento(pasta, linhas_arquivo, itens_calc, compras, resultados):
    """Reescreve o arquivo: mantém as regras do usuário/análise, registra as decisões automáticas (fonte "automatico", com o
    dispositivo) e lista os produtos com redução ainda sem enquadramento, para análise."""
    existentes = {(str(r.get('tipo', '')).upper(), _dig(r.get('codigo'))) for r in linhas_arquivo}
    novas = []
    for (pos, e, c, item, linhas, _), res in zip(itens_calc, resultados):
        for l, r in zip(linhas, res):
            cm = compras[l['linha_idx']]
            if not _tem_reducao(cm):
                continue
            auto, editado = cm.get('_enq_auto'), cm.get('_enq_editado')
            if not auto and not editado and r['status'] != 'dados_fiscais_pendentes':
                continue                                             # decidido por regra do usuário
            chave = _chave_enq(cm)
            if chave in existentes:
                continue
            existentes.add(chave)
            if editado:
                novas.append(dict(editado, fonte='usuario', justificativa=(str(editado.get('justificativa') or '') +
                                                                            ' [editado pelo usuário sobre a decisão automática]').strip()))
                continue
            novas.append(dict(tipo=chave[0], codigo=chave[1], descricao=cm.get('descricao', ''), ncm=cm.get('ncm', ''), cest=cm.get('cest', ''),
                              p_red_bc=f"{l['p_red_efetivo']:g}", origem_reducao=l.get('p_red_origem', ''),
                              reducao=ROTULO_ENQ.get(auto['reducao'], '') if auto else '', dispositivo=auto['dispositivo'] if auto else '',
                              evidencia_tipo='descricao_principio_ativo' if auto else '',
                              evidencia_fonte=f"descrição da NF-e {cm.get('chave_nfe', '')} item {cm.get('n_item', '')}: {cm.get('descricao', '')}" if auto else '',
                              apresentacao_confirmada='sim' if auto else '', vigencia_condicoes='sim' if auto else '',
                              justificativa=auto['justificativa'] if auto else '', fonte=FONTE_AUTO if auto else '',
                              sugestao='' if auto else EA.sugestao_marca(cm.get('descricao', ''))))
    if not novas and not linhas_arquivo:
        return None
    caminho = os.path.join(pasta, ENQUADRAMENTO)
    with open(caminho, 'w', newline='', encoding='utf-8-sig') as f:
        w = csv.DictWriter(f, fieldnames=CAMPOS_ENQ, delimiter=';', extrasaction='ignore')
        w.writeheader()
        w.writerows([{k: r.get(k, '') for k in CAMPOS_ENQ} for r in linhas_arquivo] + novas)
    return caminho


def _atualizar_leiame(pasta, cred, n_enq_pend):
    """Troca a linha "MODELO VAZIO" do LEIAME.txt pela situação do levantamento."""
    caminho = os.path.join(pasta, 'LEIAME.txt')
    if not os.path.exists(caminho):
        return
    with open(caminho, encoding='utf-8') as f:
        linhas = f.read().splitlines()
    novas = [f"resultado_levantamento.xlsx    levantamento concluído em {date.today().strftime('%d/%m/%Y')}: crédito {_brl(cred['total'])} (Anexo {cred['anexo']})",
             'relatorio_parser/trilha_levantamento.xlsx  trilha técnica (triagem, localização, alocação, fora do crédito, uso de notas)',
             'relatorio_levantamento.pdf     relatório (python scripts/gerar_relatorio.py <pasta do cliente>)']
    if n_enq_pend:
        novas.append(f'{ENQUADRAMENTO}      {n_enq_pend} produto(s) com redução de BC sem enquadramento: ficam sem valor até a análise (evidência na própria planilha) e nova rodada')
    saida, trocou = [], False
    for t in linhas:
        if t.startswith(('resultado_levantamento.xlsx', 'relatorio_parser/trilha_levantamento', 'relatorio_levantamento.pdf', ENQUADRAMENTO)):
            if not trocou:
                saida += novas
                trocou = True
            continue
        saida.append(t)
    if not trocou:
        saida += [''] + novas
    with open(caminho, 'w', encoding='utf-8') as f:
        f.write('\n'.join(saida) + '\n')


def _aplicar_entradas(pasta, compras):
    """<cliente>/entradas.csv (chave;data_entrada;origem): registro de entrada de cada nota (EFD C100 DT_E_S, livro de entradas,
    manifestação do destinatário). É a evidência da entrada; nota com entrada registrada na vigência deixa de suprir o estoque."""
    caminho = os.path.join(pasta, ENTRADAS)
    if not os.path.exists(caminho):
        return 0
    reg = {}
    for r in _ler_csv(caminho):
        ch = _dig(r.get('chave') or r.get('chave_nfe'))
        if len(ch) == 44 and r.get('data_entrada'):
            dt = str(r['data_entrada']).strip()
            if '/' in dt:
                d, m, a = dt[:10].split('/')
                dt = f'{a}-{m}-{d}'
            reg[ch] = (dt[:10], str(r.get('origem') or 'informado'))
    n = 0
    for c in compras:
        if c.get('chave_nfe') in reg:
            c['_data_entrada'], c['_origem_entrada'] = reg[c['chave_nfe']]
            n += 1
    return n


def _incorporar_complementares(compras):
    """NF complementar (finNFe 2) com refNFe: soma a base de ST, o ICMS-ST, o FCP e o valor complementados ao item da nota
    original (mesmo código de produto, senão mesmo EAN, senão mesmo número de item). Sem item correspondente, as linhas da
    original levam observação. A complementar não supre estoque (alocacao.elegivel)."""
    por_chave = defaultdict(list)
    for c in compras:
        por_chave[c.get('chave_nfe')].append(c)
    campos = ('vbc_st', 'vbc_st_ret', 'vicms_st', 'vicms_st_ret', 'vfcp_st', 'vfcp_st_ret', 'vbc_fcp_st', 'vbc_fcp_st_ret', 'vl_merc',
              'frete', 'seguro', 'outras_despesas', 'desconto', 'vicms', 'vicms_substituto')
    incorporadas = sem_vinculo = 0
    for cp in compras:
        if str(cp.get('fin_nfe') or '') != '2':
            continue
        for ref in [x for x in str(cp.get('ref_nfe') or '').split('|') if x]:
            originais = por_chave.get(ref, [])
            alvo = (next((o for o in originais if o.get('cod_produto') and o.get('cod_produto') == cp.get('cod_produto')), None)
                    or next((o for o in originais if L.gtin14(o.get('ean')) and L.gtin14(o.get('ean')) == L.gtin14(cp.get('ean'))), None)
                    or next((o for o in originais if str(o.get('n_item')) == str(cp.get('n_item'))), None))
            if alvo:
                for k in campos:
                    alvo[k] = float(alvo.get(k) or 0) + float(cp.get(k) or 0)
                alvo.setdefault('_complementada', []).append(f"{cp.get('chave_nfe')} item {cp.get('n_item')}")
                incorporadas += 1
                break
        else:
            sem_vinculo += 1
            for ref in [x for x in str(cp.get('ref_nfe') or '').split('|') if x]:
                for o in por_chave.get(ref, []):
                    o['_tem_complementar'] = True
    return dict(incorporadas=incorporadas, sem_vinculo=sem_vinculo)


def _gravar_identidade(pasta, linhas_arquivo, pendentes):
    """Reescreve identidade_cat68.csv: preserva as decisões preenchidas e lista as mercadorias sem identidade confirmada."""
    existentes = {(L.gtin14(r.get('ean')), str(r.get('anexo')).strip(), str(r.get('item')).strip()) for r in linhas_arquivo}
    novas = []
    for r in pendentes:
        k = (L.gtin14(r['ean']), str(r['anexo']).strip(), str(r['item']).strip())
        if k in existentes:
            continue
        existentes.add(k)
        novas.append(r)
    if not novas and not linhas_arquivo:
        return None
    caminho = os.path.join(pasta, IDENTIDADE)
    with open(caminho, 'w', newline='', encoding='utf-8-sig') as f:
        w = csv.DictWriter(f, fieldnames=CAMPOS_ID, delimiter=';', extrasaction='ignore')
        w.writeheader()
        w.writerows([{k: r.get(k, '') for k in CAMPOS_ID} for r in linhas_arquivo + novas])
    return caminho


def _rank_evidencia(loc_res) -> int:
    """Quem tem a prova mais forte escolhe primeiro as notas: 0 = EAN; 1 = descrição A/B sem ambiguidade; 2 = o resto."""
    if loc_res['metodo'] in ('ean', 'ean_trib', 'ean_embalagem'):
        return 0
    if loc_res['metodo'] == 'descricao' and loc_res['nivel'] in ('A', 'B') and not loc_res.get('ambiguo'):
        return 1
    return 2


def levantar(pasta: str, progresso=None) -> dict:
    compras, campos_parser, estoque, cli, resumo = carregar(pasta)
    if cli.get('regime') not in ('SN', 'RPA'):
        raise ValueError(f"Regime do detentor do estoque inválido ({cli.get('regime')!r}). Informe SN (Anexo V) ou RPA (Anexo IV) no "
                         'cliente.json ou rode preparar_cliente.py com --regime SN|RPA. Excesso de sublimite: decidir com o usuário (ICMS pelo regime normal = RPA).')
    n_entradas = _aplicar_entradas(pasta, compras)
    compl = _incorporar_complementares(compras)
    linhas_ident, regras_ident = _carregar_identidade(pasta)
    linhas_enq, regras_enq, auto_ant = _carregar_enquadramento(pasta)
    _aplicar_enquadramento(compras, regras_enq, auto_ant, cli.get('regime', ''))
    tri, loc = T.Triagem(), L.Localizador(compras)
    por_data = defaultdict(list)
    for e in estoque:
        por_data[_data(e['data_posicao'])].append(e)
    triagem_rows, pend = [], []
    trabalho = []                            # itens triados, na ordem de apresentação
    vistos_pend = set()
    sem_posicao = defaultdict(dict)        # data de revogação -> {EAN: valor em estoque} (sem duplicar entre posições)

    def pendencia(tipo, sev, pos, e, detalhe, acao, unico=False):
        chave = (tipo, e.get('ean') if e else None, detalhe if unico else None)
        if unico and chave in vistos_pend:
            return
        vistos_pend.add(chave)
        pend.append([tipo, sev, _fmt(pos), (e or {}).get('ean', ''), (e or {}).get('descricao', ''),
                     round((e or {}).get('valor_total', 0.0), 2) if e else None, detalhe, acao])

    todas_posicoes = set(por_data)
    for pos in sorted(por_data):
        itens = por_data[pos]
        classificados = []
        for e in itens:
            c = tri.classificar(e['ncm'], e.get('cest', ''), e.get('descricao', ''), e.get('cest_origem', ''))
            classificados.append((e, c))
            if c['status'] == 'ambiguo':
                alt = '; '.join(f'Anexo {a} em {_fmt(d)}' for a, d in c['alternativas'])
                pendencia('Item em mais de um anexo', 'Conferir', pos, e, f"{'; '.join(c['flags'])}. Alternativas: {alt}",
                          'Definir o anexo (e a data de revogação) com o CEST correto do produto', unico=True)
            elif c['status'] == 'revisar_parcial_sem_cest':
                alt = '; '.join(f'Anexo {a} em {_fmt(d)}' for a, d in c['alternativas'])
                pendencia('Anexo parcial sem CEST', 'Conferir', pos, e, f"{'; '.join(c['flags'])}. Possível: {alt}",
                          'Obter o CEST do produto (NF de compra) e confirmar NCM+CEST', unico=True)
        part = T.particionar([dict(c, **{'_e': e}) for e, c in classificados], todas_posicoes)   # sem_posicao = datas SEM arquivo
        for dt, lst in part['sem_posicao'].items():
            for it in lst:
                sem_posicao[dt].setdefault(it['_e']['ean'], it['_e']['valor_total'])
        for it in part['por_posicao'][pos]:
            e, c = it['_e'], it
            item = dict(ean=e['ean'], descricao=e['descricao'], ncm=e['ncm'], qtd=e['qtd'], valor=e['valor'],
                        aliq_totalizador=e['aliq_totalizador'])
            triagem_rows.append([_fmt(pos), e['ean'], e['descricao'], e['ncm'], e['qtd'], e['valor'], round(e['valor_total'], 2),
                                 c['anexo'], c['segmento'], c['item'], c['criterio'], 'sim' if c['cest_confirma'] else '', _fmt(c['data_revogacao']), e.get('cest', ''),
                                 e.get('cest_origem', ''), c['ato'], '; '.join(c['flags'])])
            if e.get('cest_origem') in ('nfe_ean_ambiguo',) and c['criterio'] == 'NCM+CEST':
                pendencia('CEST ambíguo', 'Conferir', pos, e, 'Mais de um CEST nas notas para este EAN (usado o mais frequente)',
                          'Confirmar o CEST do produto', unico=True)
            trabalho.append(dict(ordem=len(trabalho), pos=pos, e=e, c=c, item=item))
        if progresso:
            progresso(pos, len(part['por_posicao'][pos]))

    # --- alocação: UM livro-razão de consumo para todo o levantamento (uso único de cada item de nota) -------------------
    consumo = {}
    for t in trabalho:
        t['primeiro'] = loc.localizar(t['item'])
        t['rank'] = _rank_evidencia(t['primeiro'])
    fila = sorted(trabalho, key=lambda t: (t['rank'], t['ordem']))       # prova mais forte escolhe as notas primeiro
    for t in fila:
        t['estrito'] = _alocar_item(t['item'], t['c'], t['primeiro'], loc, compras, consumo, exigir_prova=True)
    A.verificar_consumo(consumo, compras)                                   # trava: nenhum item de nota além da quantidade comprada
    uso_estrito = dict(consumo)
    for t in fila:                                                          # simulação do que ficou de fora, só com a sobra das notas
        t['sombra'] = None
        if t['estrito']['restante'] > A.EPS:
            sombra = _alocar_item(t['item'], t['c'], t['primeiro'], loc, compras, consumo, exigir_prova=False,
                                  restante=t['estrito']['restante'])
            t['sombra'] = sombra if sombra['linhas'] else None
    A.verificar_consumo(consumo, compras)

    loc_rows, aloc_rows, itens_calc, sombra_itens = [], [], [], []
    for t in trabalho:                                                      # volta à ordem de apresentação
        loc_rows.append(_registrar_item(t, aloc_rows, pendencia, itens_calc, sombra_itens))

    for dt, eans in sem_posicao.items():
        n, v = len(eans), sum(eans.values())
        pend.insert(0, ['Falta posição de estoque', 'Alerta', '', '', f'{n} produto(s) de anexo revogado em {_fmt(dt)}, sem posição enviada para essa data',
                        None, f'Revogação em {_fmt(dt)}: o crédito usa a posição de {_fmt(dt - timedelta(days=1))}, que não veio no pacote; esses itens ficam sem crédito',
                        f'Se houver estoque em {_fmt(dt - timedelta(days=1))}, enviar a posição dessa data'])

    if compl['incorporadas']:
        pend.append(['NF complementar incorporada', 'Info', '', '', f"{compl['incorporadas']} item(ns) de NF complementar somados ao item original", None,
                     'A base de ST, o ICMS-ST e o valor complementados entram no item da nota original (art. 4º, II)', 'Nenhuma'])
    if compl['sem_vinculo']:
        pend.append(['NF complementar sem vínculo ao item', 'Conferir', '', '', f"{compl['sem_vinculo']} item(ns) de NF complementar sem item correspondente na original",
                     None, 'Não foi possível ligar o complemento a um item da nota referenciada (refNFe); as linhas dessa nota levam observação',
                     'Conferir a nota complementar e o item da original'])
    if not n_entradas:
        pend.append(['Entrada sem registro', 'Info', '', '', 'Nenhum registro de entrada informado (entradas.csv)', None,
                     'A entrada de cada nota fica evidenciada só pela emissão e pelo dhSaiEnt do emitente, que não comprova o recebimento',
                     f'Para evidenciar a entrada, informar {ENTRADAS} (chave;data_entrada;origem) a partir da EFD, do livro de entradas ou da manifestação'])
    cred = _aplicar_credito(itens_calc, cli['regime'], pend, regras_ident)
    cred['arquivo_identidade'] = _gravar_identidade(pasta, linhas_ident, cred['ident_pend'])
    cred.update(_fora_do_credito(sombra_itens, cli['regime'], pend))
    cred['uso_notas'] = _uso_notas(itens_calc, compras, uso_estrito)
    arq_enq = _gravar_enquadramento(pasta, linhas_enq, itens_calc, compras, cred['resultados'])
    cred['enq_pendentes'] = sum(1 for r in (_ler_csv(arq_enq) if arq_enq else []) if not str(r.get('reducao') or '').strip())
    linhas_finais = _ler_csv(arq_enq) if arq_enq else []
    cred['enq_aplicados'] = sum(1 for r in linhas_finais if str(r.get('reducao') or '').strip().lower() in VALORES_ENQ)
    cred['enq_automaticos'] = sum(1 for r in linhas_finais if str(r.get('fonte') or '') == FONTE_AUTO)
    caminho = _escrever(pasta, cli, campos_parser, compras, triagem_rows, loc_rows, aloc_rows, pend, resumo, cred)
    _atualizar_leiame(pasta, cred, cred['enq_pendentes'])
    return dict(arquivo=caminho, triados=len(triagem_rows), alocacoes=len(aloc_rows), pendencias=len(pend),
                por_situacao=dict(Counter(r[16] for r in loc_rows)),
                valor_triado=round(sum(r[6] for r in triagem_rows), 2),
                valor_coberto=round(sum(r[13] for r in loc_rows), 2), credito_total=cred['total'],
                credito_interpretativo=cred['interpretativo'], mercadorias_definitivas=cred['n_definitivo'],
                mercadorias_interpretativas=cred['n_interpretativo'],
                fora_do_credito=cred['fora_total'], itens_fora_do_credito=len(cred['fora_itens']),
                notas_compartilhadas=len(cred['uso_notas']), anexo=cred['anexo'])


def _acao_pendente(alertas):
    """Ação sugerida conforme a causa da linha não calculada."""
    acoes = []
    for a in alertas:
        if a.startswith('Redução de BC'):
            acoes.append(f'Preencher a coluna reducao (aplicavel ou nao_aplicavel) em {ENQUADRAMENTO} e rodar o levantamento de novo')
        elif a.startswith('alíquota interna ausente'):
            acoes.append('Informar a alíquota interna do produto (a nota não traz a da ST e o cadastro do estoque está vazio ou 0%)')
        elif 'CST/CRT' in a or a.startswith('Regime'):
            acoes.append('Conferir o CST/CSOSN e o CRT do fornecedor na nota')
        elif a.startswith('UF'):
            acoes.append('Conferir as UFs do emitente e do destinatário na nota')
        elif a.startswith('combinação não prevista'):
            acoes.append('Levar ao usuário: combinação sem fórmula na portaria')
        else:
            acoes.append('Completar o dado fiscal indicado no detalhe')
    return '; '.join(dict.fromkeys(acoes))


def _mes_seguinte(d):
    return date(d.year + (d.month == 12), d.month % 12 + 1, 1)


def _carregar_identidade(pasta):
    """<cliente>/identidade_cat68.csv: decisões da análise/usuário sobre a identidade da mercadoria no item da CAT 68.
    Devolve (linhas preenchidas a preservar, {(EAN, anexo, item): linha})."""
    caminho = os.path.join(pasta, IDENTIDADE)
    if not os.path.exists(caminho):
        return [], {}
    linhas, regras = [], {}
    for r in _ler_csv(caminho):
        dec = str(r.get('confirmado') or '').strip().lower()
        if dec in ('sim', 'nao', 'não') or str(r.get('justificativa') or '').strip() or str(r.get('evidencia_fonte') or '').strip():
            linhas.append(r)
            if dec in ('sim', 'nao', 'não'):
                regras[(L.gtin14(r.get('ean')), str(r.get('anexo')).strip(), str(r.get('item')).strip())] = r
    return linhas, regras


def _triagem_confirmada(e, c, regras_ident=None):
    """Identidade da mercadoria no item da CAT 68 (identidade_cat68.avaliar) ou decisão registrada no arquivo de identidade.
    Devolve (situacao, motivo, avaliacao): situacao 'confirmada' | 'nao_confirmada' | 'excluida'."""
    av = ID.avaliar(dict(descricao=e.get('descricao', ''), ncm=e.get('ncm', ''), cest=e.get('cest', ''), cest_origem=e.get('cest_origem', '')),
                    c.get('anexo', ''), c.get('item', ''))
    regra = (regras_ident or {}).get((L.gtin14(e.get('ean')), str(c.get('anexo', '')).strip(), str(c.get('item', '')).strip()))
    if regra:
        dec = str(regra.get('confirmado') or '').strip().lower()
        if dec in ('nao', 'não'):
            return 'excluida', 'a análise concluiu que a mercadoria não corresponde ao item da CAT 68: ' + str(regra.get('justificativa') or ''), av
        ok, falta = _evidencia_ok(regra, exige_dispositivo=False)
        if ok:
            av['evidencias'] = av['evidencias'] + [f"{regra.get('evidencia_tipo')}: {regra.get('evidencia_fonte')}"]
            return 'confirmada', '', av
        return 'nao_confirmada', f'identidade confirmada sem evidência completa (falta {falta})', av
    if av['confirmada']:
        return 'confirmada', '', av
    return 'nao_confirmada', av['motivo'], av


def _aplicar_credito(itens_calc, regime, pend, regras_ident=None):
    """Aplica a fórmula da CAT 28/2020 (Anexo V para SN, IV para RPA) linha a linha e consolida por mercadoria. Todo crédito
    calculado entra no total unificado (e no lançamento); as premissas adotadas (motivos) ficam como observações do cálculo.
    Sem valor: mercadoria excluída pela análise de identidade, sem nota, ou pendente (dado fiscal ausente). No retorno,
    `n_definitivo` = mercadorias com crédito, `interpretativo`/`n_interpretativo` = parte do total que leva observação."""
    anexo2, resumo_prod, resultados = [], [], []
    por_data = defaultdict(lambda: dict(total=0.0, itens=0))
    tot = dict(total=0.0, alt_vprod=0.0, interp=0.0, soma_linhas=0.0, ajuste_piso=0.0, n_def=0, n_interp=0)
    por_motivo = defaultdict(lambda: [0, 0.0])
    ident_pend = []
    for pos, e, c, item, linhas, descoberto in itens_calc:
        res = [CR.calcular_linha(regime, l) for l in linhas]
        resultados.append(res)
        cons = CR.consolidar_item(linhas, res)
        tri_sit, tri_motivo, tri_av = _triagem_confirmada(e, c, regras_ident)
        tri_ok = tri_sit == 'confirmada'
        if linhas and tri_sit != 'confirmada':
            ident_pend.append(dict(ean=e['ean'], descricao=e['descricao'], ncm=e['ncm'], cest=e.get('cest', ''), anexo=c.get('anexo', ''),
                                   item=c.get('item', ''), descricao_legal=tri_av.get('descricao_legal', ''), motivo=tri_motivo,
                                   sugestao='; '.join(tri_av.get('evidencias', []))))
        motivos = list(cons['motivos']) + ([tri_motivo] if (tri_motivo and linhas and tri_sit == 'nao_confirmada') else [])
        # todo crédito calculado entra no total unificado; as ressalvas (motivos) viajam como observação do cálculo
        com_ressalva = bool(linhas) and (not cons['definitivo'] or not tri_ok)
        definitivo = bool(linhas) and tri_sit != 'excluida' and not (cons['pendente'] and cons['credito'] == 0)
        for l, r in zip(linhas, res):
            alts = '; '.join(f'{k}: {_brl(v)}' for k, v in (r.get('alternativas') or {}).items())
            anexo2.append([e.get('cod_produto', ''), e['ncm'], e.get('cest', ''), item['qtd'], e.get('unidade', ''), r.get('aliquota_usada'),
                           l['chave_nfe'], l['n_item'], l['qtd_nf'], l['unid_nf'], l['fator'], round(l['qtd_usada_estoque'] / l['fator'], 4),
                           round(l['vl_liq_usado'], 2), round(l['bc_st_usada'], 2), r.get('p_red_usado') or None,
                           r['credito'], r['formula'], r['linha_tabela'], '; '.join(r['alertas']), e['ean'], e['descricao'],
                           _fmt(c['data_revogacao']), round(l['vl_merc_usado'], 2), round(l['desconto_usado'], 2) or None, r['alt_vprod'],
                           '; '.join(r.get('motivos') or []), r.get('p_red_origem_usado', ''),
                           l.get('p_red_bc_st') or None, l.get('p_red_bc') or None, r.get('enquadramento', ''), alts,
                           l.get('aliquota_origem', ''), l.get('evidencia_entrada', ''), 'sim' if l.get('revisar_entrada') else '',
                           '; '.join(tri_av.get('evidencias', []))])
        q = item['qtd'] or 1
        credito_def = cons['credito'] if definitivo else 0.0
        credito_int = cons['credito'] if (com_ressalva and definitivo) else 0.0       # parte do total com observação (só informativo)
        if not linhas:
            situacao = 'Sem crédito: sem prova documental' if item.get('sem_prova') else 'Sem crédito: sem nota anterior à vigência'
        elif tri_sit == 'excluida':
            situacao = 'Sem crédito: a mercadoria não corresponde ao item da CAT 68 (análise)'
        elif cons['pendente'] and cons['credito'] == 0:
            situacao = 'Pendente: dados fiscais ausentes ou combinação não prevista na portaria'
        elif cons['credito'] <= 0:
            situacao = 'Sem crédito: sem base de ST' if cons['sem_base'] == cons['nlinhas'] else 'Sem crédito: BC ST não supera o valor da mercadoria'
        else:
            situacao = 'Crédito com observações' if com_ressalva else 'Crédito'
        if cons['pendente']:
            alertas_p = [a for r in res if r['credito'] is None for a in r['alertas']]
            pend.append(['Crédito não calculado', 'Alerta', _fmt(pos), e['ean'], e['descricao'], round(e['valor_total'], 2),
                         '; '.join(dict.fromkeys(alertas_p)), _acao_pendente(alertas_p)])
        if com_ressalva and definitivo:
            for m in motivos:
                por_motivo[m][0] += 1
                por_motivo[m][1] += credito_int
        resumo_prod.append([e.get('cod_produto', ''), e['ean'], e['descricao'], _fmt(c['data_revogacao']), item['qtd'], round(sum(l['vl_liq_usado'] for l in linhas), 2),
                            round(cons['total_bc'], 2), credito_def, round(cons['total_vl'] / q, 4), round(cons['total_bc'] / q, 4),
                            round(credito_def / q, 4), round(descoberto, 4), situacao,
                            '; '.join(motivos)])
        if definitivo:
            d = por_data[c['data_revogacao']]
            d['total'] += credito_def
            d['itens'] += 1 if credito_def > 0 else 0
            tot['total'] += credito_def
            tot['soma_linhas'] += cons['total_literal']
            tot['ajuste_piso'] += cons['ajuste_piso']
            tot['n_def'] += 1 if credito_def > 0 else 0
            tot['alt_vprod'] += max(0.0, sum((r['alt_vprod'] if r['alt_vprod'] is not None else (r['credito'] or 0.0)) for r in res))
        if definitivo and com_ressalva:
            tot['interp'] += credito_int                       # quanto do total tem observação (informativo)
            tot['n_interp'] += 1 if credito_int > 0 else 0
    for m, (n, v) in sorted(por_motivo.items(), key=lambda kv: -kv[1][1]):
        pend.append(['Observação do cálculo', 'Conferir', '', '', f'{n} mercadoria(s): {m}', round(v, 2),
                     'Crédito calculado e somado ao total; a observação indica a premissa adotada e o documento ou dado que a confirma',
                     'Guardar o documento ou cadastro que sustenta a premissa (dispositivo, cadastro fiscal, nota)'])
    comp = []
    for dt, d in sorted(por_data.items()):
        total = round(d['total'], 2)
        if total <= 0:
            continue
        obs = f"{d['itens']} mercadorias com crédito"
        if regime == 'SN':
            comp.append([_mes_seguinte(dt).strftime('%m/%Y'), 'única', total, 'Simples Nacional',
                         'Dedução do ICMS devido no PGDAS-D (campo "redução da base de cálculo"), no mês posterior ao da exclusão',
                         obs + '. O que exceder o ICMS devido do mês compensa nos meses seguintes.'])
        else:
            mes = date(dt.year, dt.month, 1)
            parcela = int(total * 100 / 12) / 100                    # arredonda para baixo: a última nunca fica negativa
            for k in range(12):
                m = date(mes.year + (mes.month - 1 + k) // 12, (mes.month - 1 + k) % 12 + 1, 1)
                valor = parcela if k < 11 else round(total - 11 * parcela, 2)      # a última absorve o arredondamento
                comp.append([m.strftime('%m/%Y'), f'{k + 1}/12', valor, 'RPA',
                             'EFD Bloco E, código SP020750, "Outros Créditos", com menção à Portaria CAT 28/2020', obs if k == 0 else ''])
    return dict(anexo2=anexo2, resumo_prod=resumo_prod, comp=comp, resultados=resultados,
                total=round(tot['total'], 2), alt_vprod=round(tot['alt_vprod'], 2), interpretativo=round(tot['interp'], 2),
                n_definitivo=tot['n_def'], n_interpretativo=tot['n_interp'], soma_linhas=round(tot['soma_linhas'], 2),
                ajuste_piso=round(tot['ajuste_piso'], 2), motivos_interp={m: (n, round(v, 2)) for m, (n, v) in por_motivo.items()},
                anexo='V' if regime == 'SN' else 'IV', ident_pend=ident_pend)


def _alocar_item(item, c, primeiro, loc, compras, consumo, exigir_prova, restante=None):
    """Seleciona as notas que suprem o item (regras de ouro 1 a 4). Com exigir_prova só entra o que está provado."""
    data_rev = c['data_revogacao']
    grupos_usados = set()
    cache_fator = {}

    def fator_fn(it, linha):
        chave = (it['ean'], id(linha))
        if chave not in cache_fator:
            cache_fator[chave] = F.fator_linha(it, linha)
        return cache_fator[chave]

    falta = item['qtd'] if restante is None else restante
    total = dict(linhas=[], descartadas={})

    candidatos = []
    def rodar(grupo):
        candidatos.append(grupo)

    metodo, nivel = primeiro['metodo'], primeiro['nivel']
    obs = list(primeiro.get('observacoes', []))
    if metodo in ('ean', 'ean_trib', 'ean_embalagem'):
        grupos_usados |= {loc.chave_produto[i] for i in primeiro['linhas']}
        rodar(dict(origem=metodo, confirmado=True, linhas=primeiro['linhas']))
    elif metodo == 'descricao':
        grupos_usados.add(primeiro['chave_produto'])
        ok = nivel in ('A', 'B') and not primeiro.get('ambiguo')
        rodar(dict(origem=f'descrição {nivel}' + (' (ambíguo)' if primeiro.get('ambiguo') else ''), confirmado=ok, linhas=primeiro['linhas']))
    # regra de ouro 1: completar a quantidade com outros produtos equivalentes (mesmo item, outro EAN), se faltar
    while True:
        extra = loc.por_descricao(item, excluir=grupos_usados)
        if extra['metodo'] != 'descricao':
            break
        if extra['chave_produto'] in grupos_usados:
            break
        grupos_usados.add(extra['chave_produto'])
        ok = extra['nivel'] in ('A', 'B') and not extra.get('ambiguo')
        rodar(dict(origem=f"descrição {extra['nivel']} (produto equivalente)", confirmado=ok, linhas=extra['linhas']))
        if ok is False and not extra['linhas']:
            break
    total = A.alocar(dict(item, qtd=falta), candidatos, compras, data_rev, consumo, fator_fn, exigir_prova=exigir_prova)
    return dict(linhas=total['linhas'], descartadas=total['descartadas'], restante=total['qtd_descoberta'], obs=obs, metodo=metodo, nivel=nivel)


def _linha_aloc(pos, e, data_rev, l):
    return [_fmt(pos), e['ean'], e['descricao'], _fmt(data_rev), l['origem'], l['chave_nfe'],
            l['n_item'], l['data_emissao'], l['descricao_nf'], l['ean_nf'], l['unid_nf'], l['qtd_nf'], l['fator'],
            l['fator_origem'], l['fator_confianca'], l['fator_obs'], round(l['qtd_equivalente_linha'], 4),
            round(l['qtd_usada_estoque'], 4), round(l['vl_merc_unit'], 6), round(l['bc_st_unit'], 6),
            round(l['bc_st_ret_unit'], 6), round(l['bc_st_total_unit'], 6), round(l['bc_st_usada'], 4), l['aliquota_st'],
            l['aliquota_origem'], l['cst'], l['crt_emitente'], l['uf_emitente'], l['uf_destinatario'], l['cfop'],
            l['p_red_bc_st'], 'sim' if l['sem_base_st'] else '']


def _motivo_sem_prova(l):
    if l['fator_confianca'] == 'C':
        return 'fator de conversão sem prova: ' + (l['fator_obs'] or l['fator_origem'])
    return f"localização sem prova: {l['origem']}"


def _registrar_item(t, aloc_rows, pendencia, itens_calc, sombra_itens):
    """Linhas de saída e pendências do item a partir da alocação estrita (a sombra vai para "Fora do crédito")."""
    pos, e, c, item = t['pos'], t['e'], t['c'], t['item']
    r, sombra = t['estrito'], t['sombra']
    data_rev = c['data_revogacao']
    linhas, restante, obs, metodo, nivel = r['linhas'], r['restante'], list(r['obs']), r['metodo'], r['nivel']
    coberta = item['qtd'] - restante
    itens_calc.append((pos, e, c, item, linhas, restante))
    disponivel = sum(l['qtd_equivalente_linha'] for l in linhas)
    for l in linhas:
        aloc_rows.append(_linha_aloc(pos, e, data_rev, l))
    fatores = Counter((l['fator'], l['fator_confianca']) for l in linhas)
    (fat, conf_f) = fatores.most_common(1)[0][0] if fatores else (None, '')
    if sombra:
        sombra_itens.append((pos, e, c, item, sombra['linhas']))
    sem_prova = bool(sombra) or any(k in (A.MOTIVO_LOCALIZACAO, A.MOTIVO_FATOR) for k in r['descartadas'])
    if not linhas and metodo == 'nao_localizado' and not sombra:
        situacao = 'Não localizado'
    elif not linhas and sem_prova:
        situacao = 'Sem prova suficiente (fora do crédito)'
        item['sem_prova'] = True
    elif not linhas:
        situacao = 'Sem nota elegível antes da vigência'
    else:
        situacao = 'Confirmado' + ('' if restante <= A.EPS else ' (parcial)')
    if restante > A.EPS and linhas:
        obs.append(f'descoberto {restante:g} un.')
    if sombra:
        obs.append(f"fora do crédito (sem prova): {sum(l['qtd_usada_estoque'] for l in sombra['linhas']):g} un.")
    ref = linhas or (sombra['linhas'] if sombra else [])
    desc_nota = ref[0]['descricao_nf'] if ref else ''
    ean_nota = ref[0]['ean_nf'] if ref else ''

    if situacao == 'Não localizado':
        pendencia('Não localizado nas notas', 'Alerta', pos, e, 'Nem por EAN nem por descrição; sem nota não há crédito',
                  'Pedir XMLs anteriores a 01/2024, transferências ou outras fontes')
    elif situacao == 'Sem nota elegível antes da vigência':
        pendencia('Sem nota elegível antes da vigência', 'Alerta', pos, e, f"Descartadas: {r['descartadas']}", 'Buscar XMLs mais antigos')
    elif restante > A.EPS and linhas:
        pendencia('Estoque sem nota suficiente', 'Alerta', pos, e, f'{restante:g} de {item["qtd"]:g} un. sem nota provada anterior à vigência',
                  'Só a parte coberta gera crédito; buscar XMLs mais antigos')
    sem_aliq = [l for l in linhas if l['aliquota_st'] is None]
    if sem_aliq:
        pendencia('Alíquota ausente', 'Alerta', pos, e, f'{len(sem_aliq)} linha(s) sem alíquota da ST e sem alíquota interna no estoque', 'Informar a alíquota interna do produto')
    padrao = [l for l in linhas if l.get('aliquota_origem') == A.ORIGEM_ALIQ_PADRAO and not l['sem_base_st']]
    if padrao:
        pendencia('Alíquota interna presumida (18%)', 'Conferir', pos, e,
                  f'{len(padrao)} linha(s) sem alíquota da ST na nota e sem alíquota no cadastro do estoque (ou 0%): usada a alíquota geral de SP, 18%',
                  'Informar a alíquota legal do produto no cadastro do estoque e rodar de novo; até lá a linha leva observação do cálculo')
    sem_base = [l for l in linhas if l['sem_base_st']]
    if sem_base and len(sem_base) == len(linhas):
        pendencia('Sem base de ST nas notas', 'Alerta', pos, e, 'Nenhuma nota selecionada traz base de ST (vBCST/vBCSTRet): crédito zero (CAT 28/20, art. 4º)',
                  'Ver se o fornecedor pode emitir NF complementar; NCM amplo pode ter pegado produto que nunca teve ST')
    antec = [l for l in sem_base if l.get('uf_emitente') and l.get('uf_emitente') != l.get('uf_destinatario')]
    if antec:
        pendencia('Possível antecipação pelo adquirente', 'Conferir', pos, e,
                  f'{len(antec)} nota(s) interestadual(is) sem ST na nota ({antec[0]["uf_emitente"]}); se o ICMS-ST foi recolhido por guia, a base não está no XML',
                  'Informar a BC ST da antecipação em bc_st_antecipacao (compras.csv) para essas notas')
    if sem_base and len(sem_base) < len(linhas):
        pendencia('Parte das notas sem base de ST', 'Info', pos, e, f'{len(sem_base)} de {len(linhas)} linha(s) sem base de ST', 'Crédito só sobre as linhas com base')
    return [_fmt(pos), e['ean'], e['descricao'], item['qtd'], metodo + (f' {nivel}' if metodo == 'descricao' else ''), nivel or '',
            desc_nota, ean_nota, len(linhas), round(disponivel, 4), round(coberta, 4), round(restante, 4),
            round(coberta / item['qtd'], 4) if item['qtd'] else 0, round(coberta * item['valor'], 2), fat, conf_f, situacao, '; '.join(obs[:4])]


def _fora_do_credito(sombra_itens, regime, pend):
    """Quanto se aproveitaria com mais prova (EAN, fator, nota melhor). NÃO entra no crédito e não exige ação."""
    linhas_out, itens_out, por_motivo = [], [], defaultdict(lambda: [0, 0.0])
    for pos, e, c, item, linhas in sombra_itens:
        res = [CR.calcular_linha(regime, l) for l in linhas]
        cons = CR.consolidar_item(linhas, res)
        motivos = []
        for k, (l, r) in enumerate(zip(linhas, res)):
            m = _motivo_sem_prova(l)
            if m.split(':')[0] not in motivos:
                motivos.append(m.split(':')[0])
            linhas_out.append([_fmt(pos), e['ean'], e['descricao'], _fmt(c['data_revogacao']), m, l['chave_nfe'], l['n_item'], l['data_emissao'],
                               l['descricao_nf'], round(l['qtd_usada_estoque'], 4), l['fator'], round(l['vl_liq_usado'], 2), round(l['bc_st_usada'], 2),
                               l['aliquota_st'], r['credito'], round(cons['credito'], 2) if k == 0 else None])
        itens_out.append(dict(ean=e['ean'], descricao=e['descricao'], rev=_fmt(c['data_revogacao']), anexo=c['anexo'], motivo=' + '.join(motivos),
                              potencial=round(cons['credito'], 2), linhas=len(linhas)))
        for mo in motivos:
            por_motivo[mo][0] += 1
            por_motivo[mo][1] += e['valor_total']
    itens_out.sort(key=lambda x: -x['potencial'])
    total = round(sum(x['potencial'] for x in itens_out), 2)
    if itens_out:
        estoque_fora = sum(e['valor_total'] for _, e, _, _, _ in sombra_itens)
        detalhe = '; '.join(f'{mo}: {n}' for mo, (n, v) in sorted(por_motivo.items()))
        pend.append(['Fora do crédito por falta de prova', 'Info', '', '', f'{len(itens_out)} item(ns) ({detalhe}; um item pode ter os dois motivos)',
                     round(estoque_fora, 2),
                     f'Sem prova documental o item não gera crédito (sem validação manual). Potencial se houvesse prova: {_brl(total)}. Detalhe na trilha técnica, aba "Fora do crédito"',
                     'Opcional: mais XMLs (compras antigas), EAN no cadastro do estoque ou embalagem correta na descrição aumentam o aproveitamento'])
    return dict(fora_linhas=linhas_out, fora_itens=itens_out, fora_total=total,
                fora_por_motivo={k: (n, round(v, 2)) for k, (n, v) in por_motivo.items()})   # (itens, estoque a custo)


def _uso_notas(itens_calc, compras, uso):
    """Auditoria do uso único: itens de nota que supriram mais de um item do estoque (saldo nunca negativo)."""
    quem = defaultdict(list)
    for pos, e, c, item, linhas, _ in itens_calc:
        for l in linhas:
            quem[l['linha_idx']].append((e['ean'], e['descricao'], round(l['qtd_usada_estoque'] / l['fator'], 4)))
    saida = []
    for i, lst in quem.items():
        if len(lst) < 2:
            continue
        cm = compras[i]
        usado = round(uso.get(i, 0.0), 4)
        saida.append([cm['chave_nfe'], cm['n_item'], cm.get('descricao', ''), float(cm['qtd']), usado, round(float(cm['qtd']) - usado, 4), len(lst),
                      '; '.join(f'{d[:30]} ({q:g})' for _, d, q in lst)])
    saida.sort(key=lambda r: (r[0], int(r[1] or 0)))
    return saida


def _escrever(pasta, cli, campos_parser, compras, triagem_rows, loc_rows, aloc_rows, pend, resumo, cred):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = 'Resumo'
    sit = Counter(r[16] for r in loc_rows)
    bloq = [p for p in pend if p[1] == 'BLOQUEANTE']
    regime = cli['regime']
    extra = [f'Itens triados: {len(triagem_rows)} ({_brl(sum(r[6] for r in triagem_rows))} em estoque) | linhas de alocação: {len(aloc_rows)}',
             'Situação da localização: ' + '; '.join(f'{k}: {v}' for k, v in sit.most_common()),
             f"CRÉDITO TOTAL (Anexo {cred['anexo']}): {_brl(cred['total'])} em {cred['n_definitivo']} mercadorias. É o valor do lançamento: soma de todo o crédito "
             "calculado pela fórmula do anexo, com as premissas e observações registradas.",
             f"Composição: soma das linhas {_brl(cred['soma_linhas'])} + ajuste do piso zero por mercadoria {_brl(cred['ajuste_piso'])} "
             f"= {_brl(cred['total'])}. Com VlMerc = vProd (sem descontar o desconto da nota) o crédito seria {_brl(cred['alt_vprod'])}.",
             (f"Observações do cálculo: {_brl(cred['interpretativo'])} do total em {cred['n_interpretativo']} mercadorias têm premissa a documentar: "
              + '; '.join(f"{m} ({n})" for m, (n, v) in sorted(cred['motivos_interp'].items(), key=lambda kv: -kv[1][1])[:6])
              + ". Detalhe nas abas Anexo II (observações do cálculo) e Pendências." if cred['n_interpretativo'] else
              "Observações do cálculo: nenhuma."),
             f"Fora do crédito por falta de prova: {len(cred['fora_itens'])} itens, potencial {_brl(cred['fora_total'])} (detalhe na trilha técnica; informativo, não exige ação)",
             (f"Enquadramento da redução de BC: {cred['enq_pendentes']} produto(s) aguardando em {ENQUADRAMENTO} (sem enquadramento a linha fica sem valor); "
              f"{cred['enq_aplicados']} regra(s) já aplicada(s)" if cred['enq_pendentes'] else
              f"Enquadramento da redução de BC: sem pendência ({cred['enq_aplicados']} regra(s) aplicada(s) de {ENQUADRAMENTO})"),
             f"Uso único das notas: nenhum item de nota supre mais do que comprou (verificado). {len(cred['uso_notas'])} itens de nota dividem a quantidade "
             "entre itens do estoque (auditoria na trilha técnica)",
             "A coluna 'crédito' do Anexo II traz o valor de cada linha de nota, que pode ser negativo; o crédito de cada mercadoria é a soma das "
             "suas linhas com piso zero (art. 3º, §1º; a exclusão nunca gera débito). Por isso a soma simples da coluna difere do crédito total.",
             "Limite da entrega: esta planilha e o PDF fornecem os valores e a memória para escriturar; não comprovam que o Registro de Inventário "
             "(art. 2º, II; Bloco H) nem os lançamentos do crédito (art. 3º) foram feitos.",
             f"Identidade da mercadoria na CAT 68: {len(cred.get('ident_pend', []))} mercadoria(s) sem correspondência confirmada com a descrição legal do item "
             f"(lista para análise em {IDENTIDADE}).",
             "Entrada: " + '; '.join(f'{k}: {v}' for k, v in Counter(
                 (a[-3].split(' em ')[0].split(' 2')[0] if a[-3] else 'sem evidência') for a in cred['anexo2']).most_common()) + " (linhas de nota).",
             "PONTOS EM ABERTO (não validados juridicamente; mostrados com o efeito): piso zero por mercadoria (ajuste "
             f"{_brl(cred['ajuste_piso'])} no total); VlMerc líquido de desconto (com vProd o crédito seria {_brl(cred['alt_vprod'])}); FCP com base "
             "própria calculado em cada base; analogias do Anexo IV; a parcela do art. 3º, §4º (CAT 75/08) não é "
             "identificada nem deduzida."]
    extra += ['BLOQUEANTE: ' + (p[4] if p[0] != 'Falta posição de estoque' else f"{p[0]} ({p[6]})") for p in bloq]
    RM.leia_me(ws, cli['cliente'], cli['cnpj'], regime + f" ({cli.get('origem_regime', '')})",
               f"Status: levantamento concluído (Anexo {cred['anexo']}) sobre as posições de estoque enviadas pelo cliente.", extra)
    ws.append([])
    for t in ('Regras de ouro aplicadas: (1) anexo completo = só NCM; parcial = NCM+CEST; (2) quantidade do estoque suprida por uma ou mais notas, '
              'o máximo possível; (3) base de ST e ST retido anteriormente UNITÁRIAS por item de nota x quantidade do estoque; (4) alíquota da ST do parser, '
              'senão a alíquota interna do item no estoque; (5) só notas anteriores à vigência da exclusão (emissão < data de revogação). '
              'Anexo parcial sem CEST e NCM em mais de um anexo: o anexo e o item são escolhidos pela descrição do produto (a que mais corresponde à descrição legal).',
              (('Fórmula (Anexo V, Simples Nacional): C = (BC ST − VlMerc) x alíquota interna, por item de nota, com VlMerc líquido de desconto; '
                if regime == 'SN' else
                'Fórmulas (Anexo IV, RPA), por item de nota: fornecedor RPA ou ST retida antes = BC ST x alíquota interna; fornecedor do Simples com retenção '
                'na nota = (BC ST − VlMerc) x alíquota interna (interna) ou BC ST x alíq. interna − VlMerc x alíq. interestadual; ')
               + 'total da mercadoria = soma dos itens, com piso zero. Abas:')
              + '  Anexo II (itens 1 a 13), Resumo por produto (itens 14 a 19), Resumo por competência (lançamento), Pendências.',
              'Trilha técnica (triagem, localização, alocação de notas, fora do crédito, uso de notas): relatorio_parser/trilha_levantamento.xlsx.'):
        ws.append([t])
    larg = {'Descrição': 46, 'Descrição (estoque)': 46, 'Descrição na nota': 46, 'Produto na nota (descrição)': 46, 'Detalhe': 70,
            'Ação sugerida': 50, 'Observações': 60, 'Sinalizações': 50, 'Chave NF-e': 46, 'Segmento': 30, 'Situação': 26, 'Descrição ': 46,
            'Situação do crédito': 52, 'Alertas': 70, 'Fórmula aplicada': 46, 'Linha da tabela (Anexo IV/V)': 50, 'Forma de lançamento': 70,
            'Observação': 70, 'chave NFe': 46, 'Premissa': 110, 'Motivo': 70, 'Itens do estoque atendidos': 90}
    modelo = RM.abas(campos_parser)
    premissas = [
        [('Regime do detentor do estoque: Simples Nacional (Anexo V)' if regime == 'SN' else 'Regime do detentor do estoque: RPA, Regime Periódico de Apuração (Anexo IV)'),
         cli.get('origem_regime', '') or 'Usuário', '', 'usuário'],
        ['Regra de ouro 1: anexo de saída completa casa só por NCM; saída parcial, por NCM+CEST', 'Usuário', '30/09/2026', 'usuário'],
        ['Regras de ouro 2 a 4: notas que suprem o estoque; base de ST unitária x quantidade; alíquota da ST do parser ou do estoque; só antes da vigência', 'Usuário', '30/09/2026', 'usuário'],
        ['Posição de estoque de D vale para a revogação de D+1 (CAT 28/20, art. 2º)', 'CAT 28/2020', '', 'skill'],
        ['CEST do estoque resolvido pelo EAN nas NF-e de compra (o relatório não traz CEST)', 'skill', '', 'skill'],
        ['Fator de conversão inferido da descrição do estoque e da nota, arbitrado pelo preço; fator sem prova (confiança C) não supre o estoque', 'skill', '', 'skill'],
        ['Prova obrigatória, sem validação manual: só gera crédito a nota localizada por EAN ou por descrição A/B sem ambiguidade, com fator A/B. '
         'O resto sai do crédito automaticamente, com o motivo, na trilha técnica (aba "Fora do crédito"; medido: descrição C acerta 55% a 90%; nenhum sinal independente a separa)', 'skill', '30/09/2026', 'skill'],
        ['Uso único da nota: um livro-razão único para todas as posições; cada item de nota supre no máximo a quantidade comprada. '
         'Ordem de escolha: itens casados por EAN primeiro, depois descrição A/B; o levantamento para se a trava detectar excesso', 'Usuário', '30/09/2026', 'usuário'],
        (['Fórmula do Anexo V por item de nota: C = (BC ST − VlMerc) x alíquota interna. BC ST = vBCST + vBCSTRet unitários x quantidade usada', 'CAT 28/2020, Anexo V', '', 'skill']
         if regime == 'SN' else
         ['Fórmulas do Anexo IV por item de nota, conforme o regime do fornecedor, o responsável pela ST, a operação e a redução de BC. BC ST = vBCST + vBCSTRet unitários x quantidade usada',
          'CAT 28/2020, Anexo IV', '', 'skill']),
        *([] if regime == 'SN' else [
            ['Combinações que o Anexo IV não lista são calculadas por analogia, e entram no crédito com a observação '
             'de analogia (fundamento oficial específico a documentar): fornecedor RPA com ST retida por substituto anterior (CST 60) = BC ST x alíquota; substituto anterior ou antecipação com redução '
             '= regra RPA/RPA; antecipação interna com fornecedor do Simples = (BC ST − VlMerc) x alíquota', 'Interpretação da skill', '01/10/2026', 'fundamentar'],
            ['Alíquota interestadual do fornecedor do Simples sem pICMS: pela origem informada no XML, 4% (importado, Res. SF 13/2012) ou 12% (Res. SF 22/1989); '
             'sem origem, 12% é presunção e a linha leva a observação', 'Resoluções do Senado', '', 'skill']]),
        *([['Anexo V com redução: a fórmula sai do enquadramento jurídico (aplicável ou não ao consumidor final), pelo texto do Anexo V (art. 3º). A '
            'comparação com o ICMS-ST destacado é só diagnóstico, registrado no alerta da linha; a diferença de rótulos em relação ao Anexo IV, isoladamente, '
            'não comprova erro de redação', 'CAT 28/2020, art. 3º e Anexo V', '01/10/2026', 'skill']] if regime == 'SN' else
          [['Anexo IV, fórmula com redução: se a BC ST da nota parece já reduzida (BC ST x alíquota = ICMS-ST + ICMS próprio), o crédito fica o da fórmula e '
            'a linha leva a observação, com o valor sem reaplicar a redução como alternativa, até documentar a natureza da base', 'CAT 28/2020, Anexo IV',
            '01/10/2026', 'fundamentar']]),
        ['Identidade da mercadoria: o NCM é a busca inicial. A mercadoria é confirmada quando a descrição corresponde à descrição legal do item da CAT 68 '
         '(medicamento: NCM 3003/3004/3006 com dose ou forma; demais: termo da descrição legal) e o CEST da nota não aponta outro item; CEST ausente ou ambíguo '
         f'não impede. Sem correspondência, a análise registra a evidência em {IDENTIDADE} (GTIN/registro sanitário, documentação do fabricante...)',
         'CAT 68/2019 + CAT 28/2020', '01/10/2026', 'skill'],
        ['Documento: nota sem protocolo (cStat), com base de ST própria e retida somadas ou com NF complementar que não se liga a nenhum item fica '
         'leva observação. NF complementar ligada ao item (refNFe + código/EAN/nº do item) é somada à original (art. 4º, II). FCP com base própria: ICMS e FCP '
         'calculados cada um na sua base (a portaria usa uma alíquota única)', 'CAT 28/2020, arts. 2º e 4º', '01/10/2026', 'skill'],
        [f'Entrada: a evidência é o registro de entrada informado em {ENTRADAS} (EFD, livro de entradas, manifestação); nota com entrada registrada na '
         'vigência não supre o estoque. Sem registro, a linha registra a evidência disponível (dhSaiEnt do emitente, que não comprova o recebimento, ou só a '
         f'emissão); emissão a menos de {A.DIAS_REVISAO_ENTRADA} dias da revogação só PRIORIZA a revisão, não é critério de conformidade',
         'CAT 28/2020, art. 2º', '01/10/2026', 'skill'],
        ['Antecipação pelo adquirente: a BC ST vem do recolhimento informado (bc_st_antecipacao, coluna do compras.csv); nota interestadual sem ST e sem essa informação '
         'fica sem crédito, com a pendência "Possível antecipação pelo adquirente"', f'CAT 28/2020, Anexo {"V" if regime == "SN" else "IV"}', '', 'skill'],
        ['Art. 3º, §4º: a parcela do imposto do inciso XVI do art. 2º do RICMS (Portaria CAT 75/08) incluída na retenção não é compensável. Ela não é identificável na NF-e '
         'e a skill não a deduz: conferir se algum produto do levantamento está sujeito a ela', 'CAT 28/2020, art. 3º, §4º', '', 'confirmar'],
        ['VlMerc = valor da operação líquido de desconto (vProd + frete + seguro + outras − vDesc): nas notas CST 10 reproduz o vICMSST cobrado em 98% das linhas. Crédito por linha conforme a fórmula do anexo, sem teto pelo ICMS-ST destacado', 'Evidência nos XMLs do cliente + Anexos IV/V', '30/09/2026', 'skill'],
        ['O valor por linha pode ser negativo (BC de ST menor que o valor da mercadoria); o total da mercadoria é a soma das linhas com piso zero (nunca gera débito)', 'Interpretação da skill do art. 3º, §1º', '', 'confirmar'],
        ['Redução de BC: exige enquadramento com EVIDÊNCIA (tipo, fonte, apresentação exata, vigência e condições), qualquer que seja quem analisou: '
         'automático só para o art. 3º, XXIV com o princípio ativo escrito na descrição; análise ou usuário com GTIN/registro sanitário, documentação do '
         f'fabricante ou documento oficial, em {ENQUADRAMENTO}. pRedBc da fórmula = o legal quando o dispositivo fixa a carga; senão o pRedBCST da nota. O pRedBC '
         'da operação própria não é transportado para a ST', 'Interpretação da skill', '01/10/2026', 'skill'],
        ['Regime do fornecedor pelo código de tributação do ICMS do item: CSOSN = Simples Nacional; CST = RPA (inclui CRT 2). CRT só sem CST válido; '
         'CST e CRT em conflito: a linha leva a observação', 'Usuário', '01/10/2026', 'usuário'],
        ['Responsável pela ST pela base presente: vBCST na nota = retenção pelo fornecedor; só vBCSTRet = retenção por substituto anterior', 'Usuário', '30/09/2026', 'usuário'],
        ['Alíquota interna: pICMSST + pFCPST (CST 10/30/70); pST sozinho (CST 60/500), pois já inclui o FCP (NT 2016.002); sem alíquota na nota, a do item no estoque; '
         'sem nenhuma das duas (ou cadastro com 0%), 18% é presunção: a linha leva a observação até o cadastro fiscal informar a alíquota legal (Anexo I, item 5)', 'Usuário', '01/10/2026', 'fundamentar'],
        (['Simples Nacional: crédito deduzido do ICMS devido no PGDAS-D no mês posterior ao da exclusão; excedente compensa nos meses seguintes (art. 3º, §3º)', 'CAT 28/2020', '', 'skill']
         if regime == 'SN' else
         ['RPA: crédito lançado em 12 parcelas mensais, iguais e sucessivas, a partir do primeiro mês de vigência da exclusão, no Bloco E da EFD (código SP020750, '
          '"Outros Créditos"), com menção à portaria (art. 3º, §2º, redação da Portaria SRE 07/26)', 'CAT 28/2020', '', 'skill']),
        ['A posição de estoque enviada pelo cliente é a posição da data do nome do arquivo, tal como veio (mesmo que duas datas tenham conteúdo igual)', 'Usuário', '30/09/2026', 'usuário'],
    ]
    from openpyxl.styles import Font
    ws.append([])
    ws.append(['PREMISSAS ADOTADAS'])
    ws.cell(ws.max_row, 1).font = Font(bold=True, size=12)
    for pr in premissas:
        ws.append([f'• {pr[0]}  [{pr[1]}; {pr[3]}]'])
    dados = {'Triagem estoque': triagem_rows, 'Localização': loc_rows, 'Alocação de notas': aloc_rows,
             'Anexo II (notas selecionadas)': cred['anexo2'], 'Resumo por produto': cred['resumo_prod'],
             'Resumo por competência': cred['comp'], 'Fora do crédito': cred['fora_linhas'], 'Uso de notas': cred['uso_notas'],
             'Pendências': pend, 'Premissas': premissas}
    _gravar_abas(wb, modelo, dados, larg)
    wb.save(os.path.join(pasta, 'resultado_levantamento.xlsx'))
    trilha = Workbook()
    trilha.remove(trilha.active)
    _gravar_abas(trilha, RM.abas_trilha(), dados, larg)
    trilha.save(os.path.join(pasta, 'relatorio_parser', 'trilha_levantamento.xlsx'))
    return os.path.join(pasta, 'resultado_levantamento.xlsx')


def _gravar_abas(wb, modelo, dados, larg):
    for nome, cab in modelo.items():
        a = wb.create_sheet(nome[:31])
        a.append(cab)
        for linha in dados.get(nome, []):
            a.append(linha)
        RM.estilizar(a, cab, larg)
        if nome not in ('Premissas', 'Resumo por competência') and a.max_row > 1:
            a.auto_filter.ref = a.dimensions


if __name__ == '__main__':
    if len(sys.argv) != 2 or sys.argv[1] in ('-h', '--help'):
        print(__doc__)
        sys.exit(0 if len(sys.argv) == 2 else 1)
    if not os.path.isdir(sys.argv[1]):
        sys.exit(f'Pasta do cliente não encontrada: {sys.argv[1]} (monte com preparar_cliente.py)')
    try:
        r = levantar(sys.argv[1], lambda d, n: print(f'  posição {d}: {n} itens triados', flush=True))
    except (ValueError, FileNotFoundError) as exc:
        sys.exit(f'Erro: {exc}')
    print(json.dumps(r, ensure_ascii=False, indent=2))
